"""Mark the rows of a review sheet the AI assistant checked in a second pass (amendment of 27.10.2026).

Sets `ai_verified = 1` in every row people were to decide (`decided_by = human`)
that no person has verified yet (`verified` empty); `verified` itself stays for
people only. Readers then count these rows and record `verified_by = ai`, so a
later check by a person (verified = 1) replaces the AI's without any other change.
Works on workbooks (with --sheet) and CSV files.

With --human INITIALS (amendment "weryfikacja przez człowieka", 27.10.2026) it writes `verified = 1` instead, only
when that person said they checked every row; a person's check replaces the AI's.
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from digital_media_analysis.annotation_sheets import AI_VERIFIED, DECIDED_BY, VERIFIED, is_one


def mark(path: Path, sheet: str | None, human: bool = False) -> int:
    """Rows marked; refuses while the workbook is open in Excel."""
    column = VERIFIED if human else AI_VERIFIED
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        todo = frame[DECIDED_BY].eq("human") & frame.get(VERIFIED, pd.Series("", index=frame.index)).str.strip().ne("1")
        frame[column] = frame.get(column, pd.Series("", index=frame.index)).where(~todo, "1")
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        return int(todo.sum())
    if (path.parent / f"~${path.name}").exists():
        raise SystemExit(f"{path.name} is open in Excel; close it first.")
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1]]
    for name in [VERIFIED, column]:
        if name not in header:
            worksheet.cell(1, len(header) + 1, name)
            header.append(name)
    position = {name: header.index(name) + 1 for name in [DECIDED_BY, VERIFIED]}
    marked = 0
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row, position[DECIDED_BY]).value != "human":
            continue
        if is_one(worksheet.cell(row, position[VERIFIED]).value):
            continue
        worksheet.cell(row, header.index(column) + 1, 1)
        marked += 1
    workbook.save(path)
    return marked


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sheet-file", type=Path, required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--human", metavar="INITIALS", help="the person who checked every row (writes verified = 1)")
    args = parser.parse_args()
    who = f"person {args.human}" if args.human else "AI assistant"
    print(f"{args.sheet_file}: {mark(args.sheet_file, args.sheet, bool(args.human))} rows marked as checked by the {who}")


if __name__ == "__main__":
    main()
