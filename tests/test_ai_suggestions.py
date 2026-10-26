import pandas as pd
import pytest
from openpyxl import load_workbook

from digital_media_analysis import ai_suggestions, errors
from digital_media_analysis.annotation_sheets import read_workbook, write_workbook
from digital_media_analysis.codebook import load_codebook
from test_annotation_sheets import units
from test_errors import gold_units


def set_cells(path, sheet_name: str, values: dict[tuple[int, str], object]) -> None:
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    header = [cell.value for cell in sheet[1]]
    for (row, column), value in values.items():
        sheet.cell(row, header.index(column) + 1, value)
    workbook.save(path)


def test_source_suggestions_count_only_after_verification(tmp_path):
    path = tmp_path / "audit_A.xlsx"
    write_workbook(path, "A", "audit", units(), load_codebook())
    set_cells(path, "Reddit", {(3, "source_types"): "media"})
    suggestions = pd.DataFrame({"unit_id": ["r1", "r2"], "ai_source_types": ["institutional", "other"],
                                "ai_uncertain": [0, 1], "ai_reason": ["według GUS", "portal"]})
    assert ai_suggestions.prefill(path, "Reddit", ["unit_id"], suggestions, ["source_types"], "test AI") == 2
    sheet = load_workbook(path)["Reddit"]
    header = [cell.value for cell in sheet[1]]
    assert header[-4:] == ["ai_source_types", "ai_uncertain", "ai_reason", "verified"]
    # The person's own label stays; the empty cell gets the suggestion.
    assert [sheet.cell(row, header.index("source_types") + 1).value for row in (2, 3)] == ["institutional", "media"]

    labels, problems = read_workbook(path, load_codebook())
    assert set(labels["platform"]) == {"wikipedia"} and any("2 wierszy z podpowiedzią AI" in problem for problem in problems)
    set_cells(path, "Reddit", {(2, "verified"): 1})
    labels, _ = read_workbook(path, load_codebook(), partial=True)
    assert labels.loc[labels["platform"].eq("reddit"), "unit_id"].tolist() == ["r1"]
    with pytest.raises(ValueError, match="already has AI columns"):
        ai_suggestions.prefill(path, "Reddit", ["unit_id"], suggestions, ["source_types"], "test AI")


def test_suggested_error_causes_are_unconfirmed_until_verified(tmp_path):
    reference, predictions, unit_texts = gold_units()
    table = errors.error_table(reference, predictions, unit_texts, "rules_v2", ["media", "scientific", "other", "unknown"], "dev")
    path = tmp_path / "errors_dev.xlsx"
    errors.write_review(path, table, unit_texts)
    keys = ["unit_id", "platform", "instrument", "variable"]
    suggestions = table[keys].assign(ai_cause="inne", ai_uncertain=0, ai_reason="test")
    assert ai_suggestions.prefill(path, "Błędy", keys, suggestions, ["cause"], "test AI") == len(table)
    reviewed, problems = errors.read_review(path, table)
    assert problems == [] and reviewed["cause"].eq("").all()
    set_cells(path, "Błędy", {(2, "verified"): 1})
    reviewed, _ = errors.read_review(path, table)
    assert reviewed["cause"].tolist() == ["inne", "", "", ""]
