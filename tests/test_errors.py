import pandas as pd
import pytest
from openpyxl import load_workbook

from digital_media_analysis import errors


@pytest.mark.parametrize(("platform", "instrument", "variable", "error", "text", "link", "cause"), [
    ("reddit", "rules_v1", "cites_source", "FN", "Cats", "https://cnn.com/a", "pominięte pole url"),
    ("reddit", "gliner_v1", "cites_source", "FN", "Cats", "https://cnn.com/a", "brak w słowniku"),
    ("wikipedia", "rules_v1", "cites_source", "FP", "[[WP:SK]], drobne techniczne", "", "skrót WP:"),
    ("wikipedia", "rules_v1", "cites_source", "FP", "drobne redakcyjne, źródła/przypisy", "", "znacznik edycji"),
    ("reddit", "gliner_v1", "cites_source", "FP", "My cat", "https://i.imgur.com/x.jpg", "link niebędący źródłem"),
    # v1 finds "gus" inside "Ferguson": a substring, not a word.
    ("reddit", "rules_v1", "source_types:other", "FP", "Ferguson protests", "", "podciąg"),
    ("reddit", "rules_v1", "source_types:scientific", "FP", "Moje badanie do pracy", "", "brak kontekstu"),
    ("wikipedia", "rules_v1", "source_types:media", "FP", "https://example.com gazeta pisze", "", "wyciek kontekstu komentarza"),
])
def test_suggested_causes(platform, instrument, variable, error, text, link, cause):
    assert errors.suggest_cause(platform, instrument, variable, error, text, link) == cause
    assert cause in errors.CAUSES


def gold_units():
    reference = pd.DataFrame({
        "unit_id": ["r1", "r2", "P#1"], "platform": ["reddit", "reddit", "wikipedia"],
        "cites_source": [1, 0, 0], "edit_about_sources": pd.array([pd.NA, pd.NA, 1], dtype="Int64"),
        "source_types": ["media;scientific", "none", "none"],
    })
    predictions = pd.DataFrame({
        "unit_id": ["r1", "r2", "P#1"], "platform": ["reddit", "reddit", "wikipedia"],
        "cites_source": [1, 1, 0], "edit_about_sources": pd.array([pd.NA, pd.NA, 0], dtype="Int64"),
        "source_types": ["media;other", "unknown", "none"],
    })
    units = pd.DataFrame({
        "unit_id": ["r1", "r2", "P#1"], "platform": ["reddit", "reddit", "wikipedia"],
        "title": ["A study in the BBC", "Ask", ""], "text": ["", "Any sources?", "źródła/przypisy"], "link": ["https://bbc.com/x", "", ""],
    })
    return reference, predictions, units


def test_error_table_separates_detection_and_type_errors():
    reference, predictions, units = gold_units()
    table = errors.error_table(reference, predictions, units, "rules_v2", ["media", "scientific", "other", "unknown", "none"], "dev")
    assert table[["unit_id", "variable", "error"]].values.tolist() == [
        ["r1", "source_types:scientific", "FN"], ["r1", "source_types:other", "FP"],
        # r2: a false citation is one detection error, without type errors.
        ["r2", "cites_source", "FP"], ["P#1", "edit_about_sources", "FN"],
    ]
    assert table["split"].eq("dev").all() and table["cause"].eq("").all()


def test_review_round_trip_and_pareto(tmp_path):
    reference, predictions, units = gold_units()
    table = errors.error_table(reference, predictions, units, "rules_v2", ["media", "scientific", "other", "unknown"], "dev")
    path = tmp_path / "errors_dev.xlsx"
    errors.write_review(path, table, units)
    workbook = load_workbook(path)
    sheet = workbook["Błędy"]
    header = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, header.index("tekst") + 1).value == "A study in the BBC"
    sheet.cell(2, header.index("cause") + 1, "brak kontekstu")
    sheet.cell(3, header.index("cause") + 1, "zła przyczyna")
    workbook.save(path)
    reviewed, problems = errors.read_review(path, table)
    assert len(problems) == 1 and "zła przyczyna" in problems[0]
    assert reviewed["cause"].tolist()[0] == "brak kontekstu"

    counted = errors.pareto(reviewed.assign(cause=["brak kontekstu", "", "", ""]))
    assert counted["errors"].sum() == 4
    assert counted["cumulative_share"].iloc[-1] == pytest.approx(1.0)
    assert counted["confirmed"].sum() == 1
    errors.plot_pareto(counted, tmp_path / "pareto.png")
    assert (tmp_path / "pareto.png").exists()
