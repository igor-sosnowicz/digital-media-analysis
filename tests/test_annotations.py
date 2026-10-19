import pandas as pd
from openpyxl import load_workbook

from digital_media_analysis import annotations
from digital_media_analysis.annotation_sheets import write_workbook
from digital_media_analysis.codebook import load_codebook

MANIFEST = pd.DataFrame({
    "unit_id": ["r1", "r2", "r9", "COVID-19#3"],
    "platform": ["reddit", "reddit", "reddit", "wikipedia"],
    "split": ["pilot", "pilot", "test", "pilot"],
})


def units() -> dict[str, pd.DataFrame]:
    return {
        "reddit": pd.DataFrame({"unit_id": ["r1", "r2"], "group": "lodz", "title": "t", "text": "x", "link": ""}),
        "wikipedia": pd.DataFrame({"unit_id": ["COVID-19#3"], "group": "COVID-19", "text": "wg WHO"}),
    }


def write_filled(path, annotator="A") -> None:
    write_workbook(path, annotator, "pilot", units(), load_codebook())
    workbook = load_workbook(path)
    for sheet_name, rows in {"Reddit": ["media", "none"], "Wikipedia": ["institutional"]}.items():
        sheet = workbook[sheet_name]
        header = [cell.value for cell in sheet[1]]
        for row_number, types in enumerate(rows, start=2):
            sheet.cell(row_number, header.index("source_types") + 1, types)
            if "edit_about_sources" in header:
                sheet.cell(row_number, header.index("edit_about_sources") + 1, 0)
    workbook.save(path)


def test_import_save_and_load_round_trip(tmp_path):
    path = tmp_path / "pilot_A.xlsx"
    write_filled(path)
    labels, problems = annotations.import_workbook(path, "A", "pilot", MANIFEST, load_codebook())
    assert problems == []
    assert labels["codebook_version"].eq(load_codebook().version).all()

    annotations.save_labels(labels, annotations.label_path("pilot", "A", tmp_path))
    loaded = annotations.load_labels("pilot", tmp_path)
    assert loaded[["unit_id", "cites_source", "source_types", "edit_about_sources"]].values.tolist() == [
        ["r1", 1, "media", pd.NA],
        ["r2", 0, "none", pd.NA],
        ["COVID-19#3", 1, "institutional", 0],
    ]
    assert str(loaded["cites_source"].dtype) == "Int64"


def test_load_labels_is_empty_before_import(tmp_path):
    assert annotations.load_labels("pilot", tmp_path).empty


def test_workbook_paths_maps_annotators(tmp_path):
    for name in ["pilot_A.xlsx", "pilot_JP.xlsx", "gold_A.xlsx"]:
        (tmp_path / name).touch()
    assert list(annotations.workbook_paths("pilot", tmp_path)) == ["A", "JP"]


def test_manifest_check_reports_missing_unexpected_repeated_and_wrong_annotator():
    labels = pd.DataFrame({
        "unit_id": ["r1", "r1", "r9"],
        "platform": "reddit",
        "annotator": ["A", "A", "B"],
    })
    problems = annotations.check_against_manifest(labels, MANIFEST, "pilot", "A")
    assert any("powtórzone unit_id ['r1']" in problem for problem in problems)
    assert any("reddit: brak 1 jednostek" in problem and "r2" in problem for problem in problems)
    assert any("spoza fazy 'pilot'" in problem and "r9" in problem for problem in problems)
    assert any("wikipedia: brak 1 jednostek" in problem for problem in problems)
    assert any("annotator zawiera ['B']" in problem for problem in problems)
