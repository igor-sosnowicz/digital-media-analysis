import pandas as pd
from openpyxl import load_workbook

from digital_media_analysis import adjudication, annotations
from digital_media_analysis.codebook import load_codebook


def labels() -> pd.DataFrame:
    rows = [
        ("r1", "reddit", "A", 1, "media", pd.NA), ("r1", "reddit", "B", 1, "media", pd.NA),
        ("r2", "reddit", "A", 1, "media", pd.NA), ("r2", "reddit", "B", 0, "none", pd.NA),
        ("COVID-19#3", "wikipedia", "A", 0, "none", 1), ("COVID-19#3", "wikipedia", "B", 0, "none", 0),
        ("COVID-19#4", "wikipedia", "A", 1, "institutional", 0),
    ]
    frame = pd.DataFrame(rows, columns=["unit_id", "platform", "annotator", "cites_source", "source_types", "edit_about_sources"])
    return frame.assign(codebook_version="2.1", uncertain=0, notes=["", "", "R2?", "", "", "", ""]).astype(
        {"cites_source": "Int64", "edit_about_sources": "Int64", "uncertain": "Int64"}
    )


def units() -> dict[str, pd.DataFrame]:
    return {
        "reddit": pd.DataFrame({"unit_id": ["r1", "r2"], "group": "science", "title": "t", "text": "x", "link": ""}),
        "wikipedia": pd.DataFrame({"unit_id": ["COVID-19#3", "COVID-19#4"], "group": "COVID-19", "text": ["źródła/przypisy", "wg WHO"]}),
    }


def test_disputed_lists_each_annotators_labels():
    disputes = adjudication.disputed(labels())
    assert disputes["unit_id"].tolist() == ["r2", "COVID-19#3"]
    r2 = disputes.set_index("unit_id").loc["r2"]
    assert (r2["A: source_types"], r2["B: source_types"], r2["A: notes"]) == ("media", "none", "R2?")
    assert disputes.set_index("unit_id").loc["COVID-19#3", "A: edit_about_sources"] == "1"


def test_arbiter_round_trip_settles_the_decided_units(tmp_path):
    codebook = load_codebook()
    disputes = adjudication.disputed(labels())
    path = tmp_path / "adjudication_pilot.xlsx"
    assert adjudication.write_sheet(path, "pilot", "ARB", disputes, units(), codebook) == {"reddit": 1, "wikipedia": 1}
    workbook = load_workbook(path)
    sheet = workbook["Reddit"]
    header = [cell.value for cell in sheet[1]]
    assert [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["r2"]
    assert sheet.cell(2, header.index("B: source_types") + 1).value == "none"
    sheet.cell(2, header.index("source_types") + 1, "media")
    sheet.cell(2, header.index("notes") + 1, "R2: link do artykułu")
    workbook.save(path)

    decisions, problems = adjudication.read_decisions(path, codebook, disputes)
    assert problems == []
    assert decisions[["unit_id", "annotator", "source_types"]].values.tolist() == [["r2", "ARB", "media"]]
    adjudication.save_decisions(decisions, adjudication.decisions_path("pilot", tmp_path))
    loaded = adjudication.load_decisions("pilot", tmp_path)

    reference, conflicts = annotations.gold_reference(labels(), loaded)
    assert set(reference["unit_id"]) == {"r1", "r2", "COVID-19#4"}
    assert reference.set_index("unit_id").loc["r2", "annotator"] == "ARB"
    assert set(conflicts["unit_id"]) == {"COVID-19#3"}


def test_decisions_outside_the_disputes_are_reported(tmp_path):
    codebook = load_codebook()
    disputes = adjudication.disputed(labels())
    path = tmp_path / "adjudication_pilot.xlsx"
    adjudication.write_sheet(path, "pilot", "ARB", disputes, units(), codebook)
    _, problems = adjudication.read_decisions(path, codebook, disputes[disputes["unit_id"].eq("COVID-19#3")].iloc[0:0])
    assert problems == []
    workbook = load_workbook(path)
    sheet = workbook["Reddit"]
    header = [cell.value for cell in sheet[1]]
    sheet.cell(2, header.index("source_types") + 1, "none")
    workbook.save(path)
    _, problems = adjudication.read_decisions(path, codebook, disputes[disputes["unit_id"].eq("COVID-19#3")])
    assert any("bez sporu" in problem for problem in problems)


def test_no_decisions_leave_the_reference_unchanged(tmp_path):
    empty = adjudication.load_decisions("gold", tmp_path)
    assert empty.empty
    assert annotations.gold_reference(labels(), empty)[0].equals(annotations.gold_reference(labels())[0])
