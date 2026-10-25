import pandas as pd
import pytest
from openpyxl import load_workbook

from digital_media_analysis.annotation_sheets import coded_unit_ids, read_workbook, units_after, workbook_metadata, write_workbook
from digital_media_analysis.codebook import load_codebook


def units() -> dict[str, pd.DataFrame]:
    return {
        "reddit": pd.DataFrame({
            "unit_id": ["r1", "r2"],
            "group": "lodz",
            "title": ["Tytuł", "Inny"],
            "text": ["Według GUS…", ""],
            "link": ["", "https://wiadomosci.wp.pl/x"],
        }),
        "wikipedia": pd.DataFrame({
            "unit_id": ["COVID-19#3", "COVID-19#4"],
            "group": "COVID-19",
            "text": ["[[WP:SK]]", "wg WHO"],
        }),
    }


def fill(path, values: dict[str, list[list]]) -> None:
    workbook = load_workbook(path)
    for sheet_name, rows in values.items():
        sheet = workbook[sheet_name]
        header = [cell.value for cell in sheet[1]]
        for row_number, row_values in enumerate(rows, start=2):
            for name, value in zip(["source_types", "edit_about_sources", "uncertain", "notes"], row_values):
                if name in header:
                    sheet.cell(row_number, header.index(name) + 1, value)
    workbook.save(path)


def test_workbook_is_blind_and_prefilled(tmp_path):
    path = tmp_path / "pilot_A.xlsx"
    write_workbook(path, "A", "pilot", units(), load_codebook())
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Instrukcja", "Reddit", "Wikipedia"]
    header = [cell.value for cell in workbook["Wikipedia"][1]]
    # cites_source is derived from source_types, so annotators do not fill it in.
    assert header == ["unit_id", "strona", "opis zmiany", "source_types",
                      "edit_about_sources", "uncertain", "notes", "annotator"]
    assert "edit_about_sources" not in [cell.value for cell in workbook["Reddit"][1]]
    reddit_header = [cell.value for cell in workbook["Reddit"][1]]
    assert workbook["Reddit"].cell(2, reddit_header.index("annotator") + 1).value == "A"


def test_round_trip_returns_clean_labels_and_derives_cites_source(tmp_path):
    path = tmp_path / "pilot_A.xlsx"
    write_workbook(path, "A", "pilot", units(), load_codebook())
    fill(path, {
        "Reddit": [["Institutional; media", None, None], ["none", None, 1, "R2?"]],
        "Wikipedia": [["None", 0, None], ["institutional", "0", None]],
    })
    labels, problems = read_workbook(path, load_codebook())
    assert problems == []
    assert labels[["unit_id", "cites_source", "source_types", "uncertain"]].values.tolist() == [
        ["r1", 1, "institutional;media", 0],
        ["r2", 0, "none", 1],
        ["COVID-19#3", 0, "none", 0],
        ["COVID-19#4", 1, "institutional", 0],
    ]
    assert labels["annotator"].eq("A").all()


def test_problems_are_reported(tmp_path):
    path = tmp_path / "pilot_B.xlsx"
    write_workbook(path, "B", "pilot", units(), load_codebook())
    fill(path, {
        "Reddit": [["medai", None, None], [None, None, 2]],
        "Wikipedia": [["none;media", None, None], ["oficjalne", 1, None]],
    })
    labels, problems = read_workbook(path, load_codebook())
    assert any("r1" in problem and "nieznane kody ['medai']" in problem for problem in problems)
    assert any("r2" in problem and "source_types brak wartości" in problem for problem in problems)
    assert any("r2" in problem and "uncertain ma wartość 2" in problem for problem in problems)
    assert any("COVID-19#3" in problem and "none razem z innymi kodami" in problem for problem in problems)
    assert any("COVID-19#3" in problem and "edit_about_sources brak wartości" in problem for problem in problems)
    # Polish v1.0 codes are no longer valid.
    assert any("COVID-19#4" in problem and "nieznane kody ['oficjalne']" in problem for problem in problems)
    # A label that breaks R13 gives no cites_source.
    assert labels["cites_source"].isna().all()


def test_partial_read_skips_rows_without_labels_and_checks_the_rest(tmp_path):
    path = tmp_path / "gold_A.xlsx"
    write_workbook(path, "A", "gold", units(), load_codebook())
    fill(path, {
        "Reddit": [["media", None, None], [None, None, None]],
        "Wikipedia": [[None, None, None], ["none;media", 1, None]],
    })
    labels, problems = read_workbook(path, load_codebook(), partial=True)
    assert labels["unit_id"].tolist() == ["r1", "COVID-19#4"]
    assert problems == ["Wikipedia, wiersz 3 (COVID-19#4): none razem z innymi kodami (media;none)"]
    # Without partial the empty rows are reported.
    _, problems = read_workbook(path, load_codebook())
    assert any("r2" in problem and "source_types brak wartości" in problem for problem in problems)


def test_coded_unit_ids_lists_rows_with_any_label(tmp_path):
    path = tmp_path / "gold_A.xlsx"
    write_workbook(path, "A", "gold", units(), load_codebook())
    # A note alone counts as started work; the prefilled annotator column does not.
    fill(path, {"Reddit": [[None, None, None, "R2?"], [None, None, None]], "Wikipedia": [["none", 0, None]]})
    assert coded_unit_ids(path) == {"reddit": ["r1"], "wikipedia": ["COVID-19#3"]}
    assert workbook_metadata(path)["phase"] == "gold"


def presentation() -> pd.DataFrame:
    return pd.DataFrame({"unit_id": [f"u{i}" for i in range(1, 9)]})


def test_units_after_starts_behind_the_last_coded_unit():
    # u3 was skipped by the first annotator: it stays theirs and is not handed over.
    chosen = units_after(presentation(), ["u1", "u2", "u4"], count=3)
    assert chosen["unit_id"].tolist() == ["u5", "u6", "u7"]
    assert units_after(presentation(), ["u1", "u2", "u4"])["unit_id"].tolist() == ["u5", "u6", "u7", "u8"]


def test_units_after_without_coded_units_takes_the_first_ones():
    assert units_after(presentation(), [], count=2)["unit_id"].tolist() == ["u1", "u2"]
    assert units_after(presentation(), ["u8"], count=5).empty


def test_units_after_rejects_units_from_another_phase():
    with pytest.raises(ValueError, match="not in this phase"):
        units_after(presentation(), ["u1", "x9"])


def test_append_units_keeps_labels_and_adds_missing_rows(tmp_path):
    from digital_media_analysis.annotation_sheets import append_units

    path = tmp_path / "pilot_A.xlsx"
    first = {platform: frame.head(1) for platform, frame in units().items()}
    write_workbook(path, "A", "pilot", first, load_codebook())
    fill(path, {"Reddit": [["media", None, None]], "Wikipedia": [["none", 1, None]]})
    more = units()
    more["reddit"] = pd.concat([more["reddit"], more["reddit"].head(1).assign(unit_id="r3")], ignore_index=True)

    added = append_units(path, more, load_codebook(), counts={"reddit": 1})
    assert added == {"reddit": 1, "wikipedia": 1}
    assert coded_unit_ids(path) == {"reddit": ["r1"], "wikipedia": ["COVID-19#3"]}
    workbook = load_workbook(path)
    sheet = workbook["Reddit"]
    header = [cell.value for cell in sheet[1]]
    assert [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["r1", "r2"]
    assert sheet.cell(3, header.index("annotator") + 1).value == "A"
    assert sheet.cell(3, header.index("source_types") + 1).fill.fgColor.rgb.endswith("FFF2CC")

    fill(path, {"Reddit": [["media", None, None], ["none", None, None]], "Wikipedia": [["none", 1, None], ["institutional", 0, None]]})
    labels, problems = read_workbook(path, load_codebook())
    assert problems == []
    assert labels["unit_id"].tolist() == ["r1", "r2", "COVID-19#3", "COVID-19#4"]
    # Nothing is missing any more, so a second append adds nothing.
    assert append_units(path, units(), load_codebook()) == {"reddit": 0, "wikipedia": 0}


def test_append_units_refuses_a_workbook_open_in_excel(tmp_path):
    from digital_media_analysis.annotation_sheets import append_units

    path = tmp_path / "pilot_A.xlsx"
    write_workbook(path, "A", "pilot", units(), load_codebook())
    (tmp_path / "~$pilot_A.xlsx").write_text("lock", encoding="utf-8")
    with pytest.raises(RuntimeError, match="open in Excel"):
        append_units(path, units(), load_codebook())


def test_context_columns_are_shown_after_the_labels_and_ignored_on_read(tmp_path):
    path = tmp_path / "adjudication_pilot.xlsx"
    shown = {"reddit": units()["reddit"].assign(**{"IS: source_types": ["media", "none"]})}
    write_workbook(path, "ARB", "pilot", shown, load_codebook(), context_columns=["IS: source_types"],
                   extra_rules=["Rozstrzygnij spór."])
    workbook = load_workbook(path)
    header = [cell.value for cell in workbook["Reddit"][1]]
    assert header[-1] == "IS: source_types" and workbook["Reddit"].cell(2, len(header)).value == "media"
    assert any(row[0] == "7. Rozstrzygnij spór." for row in workbook["Instrukcja"].iter_rows(values_only=True))
    fill(path, {"Reddit": [["media", None, None], ["none", None, None]]})
    labels, problems = read_workbook(path, load_codebook())
    assert problems == [] and labels["annotator"].eq("ARB").all()
