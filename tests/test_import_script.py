import importlib.util
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

from digital_media_analysis import annotations
from digital_media_analysis.annotation_sheets import write_workbook
from digital_media_analysis.codebook import load_codebook
from digital_media_analysis.paths import PROJECT_ROOT

MANIFEST = pd.DataFrame({"unit_id": ["r1", "r2"], "platform": "reddit", "phase": "pilot", "split": "pilot"})


def load_script():
    spec = importlib.util.spec_from_file_location("import_script", PROJECT_ROOT / "scripts" / "04_import_annotations.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def workbook(path, annotator: str, codes: list[str]) -> None:
    units = {"reddit": pd.DataFrame({"unit_id": ["r1", "r2"], "group": "science", "title": "t", "text": "x", "link": ""})}
    write_workbook(path, annotator, "pilot", units, load_codebook())
    book = load_workbook(path)
    sheet = book["Reddit"]
    header = [cell.value for cell in sheet[1]]
    for row, code in enumerate(codes, start=2):
        sheet.cell(row, header.index("source_types") + 1, code)
    book.save(path)


@pytest.fixture
def script(tmp_path, monkeypatch):
    module = load_script()
    monkeypatch.setattr(module.gold, "load_manifest", lambda: MANIFEST)
    monkeypatch.setattr(module.annotations, "label_path", lambda phase, annotator: tmp_path / "labels" / f"{annotator}.csv")
    return module


def test_invalid_rows_are_left_out_and_the_rest_is_written(script, tmp_path, monkeypatch, capsys):
    workbook(tmp_path / "pilot_A.xlsx", "A", ["media", "medai"])
    monkeypatch.setattr(script.annotations, "workbook_paths", lambda phase: {"A": tmp_path / "pilot_A.xlsx"})
    monkeypatch.setattr(sys, "argv", ["04", "--phase", "pilot"])
    script.main()
    saved = annotations.read_label_file(tmp_path / "labels" / "A.csv")
    assert saved["unit_id"].tolist() == ["r1"]
    assert "1 row(s) with invalid labels left out" in capsys.readouterr().out


def test_workbook_problems_block_and_write_nothing(script, tmp_path, monkeypatch):
    # The file belongs to A, but the annotator column says B: a problem of the whole workbook.
    workbook(tmp_path / "pilot_A.xlsx", "B", ["media", "none"])
    monkeypatch.setattr(script.annotations, "workbook_paths", lambda phase: {"A": tmp_path / "pilot_A.xlsx"})
    monkeypatch.setattr(sys, "argv", ["04", "--phase", "pilot"])
    with pytest.raises(SystemExit, match="workbook problems"):
        script.main()
    assert not (tmp_path / "labels" / "A.csv").exists()
