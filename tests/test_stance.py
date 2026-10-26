import pandas as pd
from openpyxl import load_workbook

from digital_media_analysis import stance
from digital_media_analysis.codebook import load_codebook
from digital_media_analysis.instruments import stance_v1


def posts() -> pd.DataFrame:
    return pd.DataFrame({
        "unit_id": ["p1", "p2", "p3"],
        "group": "science",
        "title": ["Vaccines cause autism, wake up", "New study on bees", "My cat"],
        "text": "",
        "link": "",
    })


def fill(path, values: list[tuple]) -> None:
    workbook = load_workbook(path)
    sheet = workbook[stance.SHEET_NAME]
    header = [cell.value for cell in sheet[1]]
    for row_number, row_values in enumerate(values, start=2):
        for name, value in zip([stance.STANCE, "uncertain", "notes"], row_values):
            sheet.cell(row_number, header.index(name) + 1, value)
    workbook.save(path)


def test_stance_workbook_round_trip(tmp_path):
    path = stance.workbook_path("pilot", "A", tmp_path)
    stance.write_workbook(path, "A", "pilot", posts(), load_codebook())
    header = [cell.value for cell in load_workbook(path)[stance.SHEET_NAME][1]]
    assert header == ["unit_id", "subreddit", "tytuł", "treść", "link", "consensus_stance", "uncertain", "notes", "annotator"]
    fill(path, [("Challenges", None, None), ("neutral", 1, "R14?"), ("not_applicable", None, None)])
    labels, problems = stance.read_workbook(path, load_codebook())
    assert problems == []
    assert labels[["unit_id", "consensus_stance", "uncertain"]].values.tolist() == [
        ["p1", "challenges", 0], ["p2", "neutral", 1], ["p3", "not_applicable", 0],
    ]
    assert labels["codebook_version"].eq(load_codebook().version).all()


def test_stance_workbook_problems_and_partial(tmp_path):
    path = stance.workbook_path("gold", "B", tmp_path)
    stance.write_workbook(path, "B", "gold", posts(), load_codebook())
    fill(path, [("distrust", None, None), (None, None, None), ("supports", 2, None)])
    labels, problems = stance.read_workbook(path, load_codebook())
    assert any("p1" in problem and "spoza" in problem for problem in problems)
    assert any("p2" in problem and "brak wartości" in problem for problem in problems)
    assert any("p3" in problem and "uncertain" in problem for problem in problems)
    labels, problems = stance.read_workbook(path, load_codebook(), partial=True)
    assert "p2" not in set(labels["unit_id"])


def test_ai_suggestions_count_only_after_the_person_verifies_them(tmp_path):
    path = stance.workbook_path("gold", "C", tmp_path)
    stance.write_workbook(path, "C", "gold", posts(), load_codebook())
    fill(path, [(None, None, None), (None, None, None), ("supports", None, "own")])
    suggestions = pd.DataFrame({"unit_id": ["p1", "p2", "p3"], "ai_stance": ["challenges", "neutral", "not_applicable"],
                                "ai_uncertain": [0, 1, 0], "ai_reason": ["anti-vaccine claim", "news", "pet"]})
    assert stance.add_ai_suggestions(path, suggestions, "test AI") == 3
    workbook = load_workbook(path)
    header = [cell.value for cell in workbook[stance.SHEET_NAME][1]]
    assert header[-4:] == stance.AI_COLUMNS
    # The person's own label stays; empty ones get the suggestion.
    assert [workbook[stance.SHEET_NAME].cell(row, header.index(stance.STANCE) + 1).value for row in (2, 3, 4)] == [
        "challenges", "neutral", "supports"]
    labels, problems = stance.read_workbook(path, load_codebook())
    assert labels.empty and any("3 wierszy z podpowiedzią AI" in problem for problem in problems)
    sheet = workbook[stance.SHEET_NAME]
    sheet.cell(2, header.index("verified") + 1, 1)
    sheet.cell(4, header.index("verified") + 1, 1)
    workbook.save(path)
    labels, problems = stance.read_workbook(path, load_codebook(), partial=True)
    assert problems == [] and labels["unit_id"].tolist() == ["p1", "p3"]
    review = stance.ai_review(labels).iloc[0]
    assert (review["checked"], review["changed"]) == (2, 1)
    import pytest

    with pytest.raises(ValueError, match="already has columns"):
        stance.add_ai_suggestions(path, suggestions, "test AI")


def test_stance_labels_save_load_reference_and_agreement(tmp_path):
    rows = [("p1", "A", "challenges"), ("p2", "A", "neutral"), ("p3", "A", "supports"),
            ("p1", "B", "challenges"), ("p2", "B", "supports"), ("p3", "B", "supports")]
    for annotator in ["A", "B"]:
        part = pd.DataFrame([{"unit_id": unit, "platform": "reddit", "annotator": a, "codebook_version": "2.1",
                              "consensus_stance": value, "uncertain": 0, "notes": ""} for unit, a, value in rows if a == annotator])
        stance.save_labels(part, stance.label_path("pilot", annotator, tmp_path))
    labels = stance.load_labels("pilot", tmp_path)
    reference, conflicts = stance.reference(labels)
    assert sorted(reference["unit_id"]) == ["p1", "p3"]
    assert sorted(conflicts["unit_id"].unique()) == ["p2"]
    table = stance.agreement(labels, samples=0).set_index("variable")
    assert table.loc["consensus_stance", "units"] == 3
    assert table.loc["consensus_stance", "observed_agreement"] == 2 / 3


def test_stance_v1_decision_rule():
    scores = pd.DataFrame({
        "unit_id": ["off", "anti", "pro", "both", "flat"],
        "relevance": [0.1, 0.9, 0.9, 0.9, 0.9],
        "supports": [0.9, 0.2, 0.8, 0.7, 0.3],
        "challenges": [0.9, 0.8, 0.1, 0.75, 0.3],
    })
    labels = stance_v1.labels(scores, {"relevance": 0.5, "supports": 0.5, "challenges": 0.5})
    assert labels["consensus_stance"].tolist() == ["not_applicable", "challenges", "supports", "challenges", "neutral"]
    assert set(stance_v1.HYPOTHESES) == {"relevance", "supports", "challenges"}
    assert set(stance.stance_values(load_codebook())) == {"supports", "challenges", "neutral", "not_applicable"}
