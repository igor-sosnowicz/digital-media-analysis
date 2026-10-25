import pandas as pd
import pytest

from digital_media_analysis import h1_frame, topics
from digital_media_analysis.codebook import load_codebook
from digital_media_analysis.instruments import distrust_v1, stance_v2


def test_topics_match_whole_words_and_pick_the_main_topic():
    found = topics.load_topics()
    labels = topics.label(pd.Series({
        "a": "Vaccines cause autism? New measles outbreak", "b": "Climate change and global warming deniers",
        "c": "The evolution of smartphones", "d": "GMO foods and climate change",
    }), found)
    table = labels.set_index("unit_id")
    assert table.loc["a", "consensus_topic"] == "vaccines" and table.loc["b", "topics"] == "climate"
    # "evolution" alone is not a pattern; the theory of evolution is.
    assert table.loc["c", "consensus_topic"] == topics.NO_TOPIC
    # Ties go to the earlier topic in the file.
    assert table.loc["d", "topics"] == "climate;gmo" and table.loc["d", "consensus_topic"] == "climate"


def test_distrust_lexicon_flags_phrases():
    labels = distrust_v1.labels(pd.Series({"x": "Big Pharma paid scientists, do your own research", "y": "A new study on pharmacies"}))
    assert labels["distrust"].tolist() == [1, 0] and labels["cue"].iloc[0] == "big pharma"


def test_stance_v2_decides_from_the_strongest_topic_score():
    units = pd.DataFrame({"unit_id": ["u1", "u2", "u3"], "topics": ["vaccines;climate", "climate", ""], "text": ["a", "b", "c"]})
    assert stance_v2.pairs(units, topics.load_topics())[["unit_id", "topic"]].values.tolist() == [
        ["u1", "vaccines"], ["u1", "climate"], ["u2", "climate"]]
    scores = pd.DataFrame({"unit_id": ["u1", "u1", "u2"], "topic": ["vaccines", "climate", "climate"],
                           "supports": [0.2, 0.9, 0.3], "challenges": [0.95, 0.1, 0.3]})
    labels = stance_v2.labels(scores, units["unit_id"], {"supports": 0.5, "challenges": 0.5})
    assert labels["consensus_stance"].tolist() == ["challenges", "neutral", "not_applicable"]


def frame_tables():
    posts = pd.DataFrame({
        "subreddit": ["science", "conspiracy"], "id": ["p1", "p2"], "title": ["Vaccines", "Climate hoax"], "text": ["", "[removed]"],
        "created_utc": [1, 2], "created_at_utc": ["2010-05-01T10:00:00+00:00", "2016-03-02T10:00:00+00:00"], "url": ["", ""],
        "permalink": ["", ""], "score": [1, 1], "num_comments": [3, 0], "removed_by_category": ["", ""], "day_weight": [30.0, 31.0],
        "author": ["alice", "bob"], "topics": ["vaccines", "climate"],
    })
    comments = pd.DataFrame({
        "subreddit": ["science"], "id": ["c1"], "link_id": ["p1"], "author": ["carol"], "body": ["Vaccines are safe"],
        "created_utc": [3], "created_at_utc": ["2010-05-01T11:00:00+00:00"], "score": [2], "removed": [0], "day_weight": [30.0],
        "scanned": [40], "chosen": [20], "post_topics": ["vaccines"],
    })
    return posts, comments


def test_units_weight_comments_by_the_share_chosen_and_flag_removed_posts():
    units = h1_frame.units(*frame_tables()).set_index("unit_id")
    assert units.loc["p_p1", "weight"] == 30.0 and units.loc["c_c1", "weight"] == 60.0
    assert units.loc["p_p2", "removed"] == 1 and units.loc["c_c1", "kind"] == "comment"
    sample = h1_frame.analysis_sample(units.reset_index(), [2005, 2018], 10, seed=1)
    assert set(sample["unit_id"]) == {"p_p1", "c_c1"}


def test_panel_validation_sample_and_pseudonyms():
    sample = pd.DataFrame({"unit_id": [f"u{i}" for i in range(12)], "author": ["a"] * 4 + ["b"] * 4 + ["[deleted]"] * 4,
                           "year": [2005, 2008, 2012, 2016] * 3, "weight": 1.0})
    assert list(h1_frame.panel_authors(sample, 3)) == ["a", "b"]
    drawn = h1_frame.validation_sample(sample, [[2005, 2010], [2011, 2018]], 3, seed=2)
    assert drawn.groupby("period").size().tolist() == [3, 3] and drawn["split"].value_counts().to_dict() == {"test": 4, "dev": 2}
    names = h1_frame.pseudonym(pd.Series(["alice", "alice", ""]), "salt")
    assert names.iloc[0] == names.iloc[1] and len(names.iloc[0]) == 12 and names.iloc[2] == ""


def test_validation_workbook_round_trip(tmp_path):
    from openpyxl import load_workbook

    codebook = load_codebook()
    units = h1_frame.units(*frame_tables()).assign(year=lambda frame: frame["year"].astype(int))
    path = tmp_path / "h1v2_validation_A.xlsx"
    h1_frame.write_validation_workbook(path, "A", units, codebook)
    workbook = load_workbook(path)
    sheet = workbook["Reddit"]
    header = [cell.value for cell in sheet[1]]
    sheet.cell(2, header.index("consensus_topic") + 1, "vaccines")
    sheet.cell(2, header.index("consensus_stance") + 1, "supports")
    sheet.cell(3, header.index("consensus_topic") + 1, "weather")
    workbook.save(path)
    labels, problems = h1_frame.read_validation_workbook(path, codebook)
    assert labels.iloc[0][["consensus_topic", "consensus_stance"]].tolist() == ["vaccines", "supports"]
    assert len(problems) == 2 and any("weather" in problem for problem in problems)


def test_fetch_top_level_comments_pages_and_keeps_top_level(monkeypatch):
    pages = [[{"id": f"c{i}", "parent_id": "t3_p" if i % 2 == 0 else "t1_x", "created_utc": i} for i in range(100)],
             [{"id": "c99", "parent_id": "t1_x", "created_utc": 99}, {"id": "c100", "parent_id": "t3_p", "created_utc": 100}]]
    calls = []
    monkeypatch.setattr(h1_frame.reddit_archive, "_get", lambda client, url, params: calls.append(params) or pages[len(calls) - 1])
    found = h1_frame.fetch_top_level_comments(None, "p", scanned=1000)
    assert len(found) == 51 and found[-1]["id"] == "c100" and calls[1]["after"] == 98


@pytest.mark.parametrize("count", [0, 5])
def test_choose_draws_without_replacement(count):
    import numpy as np

    items = list(range(10))
    chosen = h1_frame.choose(items, count, np.random.default_rng(0))
    assert len(chosen) == count and chosen == sorted(set(chosen))
