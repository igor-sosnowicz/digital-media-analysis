"""Frame, analysis sample, author panel and validation sample of hypothesis H1 v2 (docs/hipotezy_v2.md).

The frame holds posts about a consensus topic (topics_v1 on title and body) in
the subreddits of `hypotheses_v2.h1`, on random days of each month (as the gold
frame), and a random sample of the top-level comments under each such post.
Unlike the gold frame it keeps the author, which the within-author analysis
needs, and removed posts with a flag, for the sensitivity analysis.

Weights: a post's weight is its day weight; a comment's is its post's day
weight times the top-level comments scanned over the comments chosen (the
frame covers the first `comments_scanned_per_post` top-level comments).
"""

import json
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from digital_media_analysis import reddit_archive
from digital_media_analysis.annotation_sheets import EXCEL_CELL_LIMIT, HEADER_FONT, LABEL_FILL, _binary
from digital_media_analysis.codebook import Codebook

COMMENTS_URL = "https://arctic-shift.photon-reddit.com/api/comments/search"
POST_COLUMNS = [*reddit_archive.COLUMNS, "author", "topics"]
COMMENT_COLUMNS = ["subreddit", "id", "link_id", "author", "body", "created_utc", "created_at_utc", "score",
                   "removed", "day_weight", "scanned", "chosen", "post_topics"]
REMOVED_TEXT = {"[removed]", "[deleted]"}
# Accounts that are not people: bots and the deleted-author placeholder.
NOT_AUTHORS = {"", "[deleted]", "automoderator", "autotldr", "remindmebot"}
UNIT_COLUMNS = ["unit_id", "kind", "author", "group", "year", "day", "introduced", "title", "text", "link",
                "day_weight", "weight", "removed", "topics"]
VALIDATION_LABELS = ["consensus_topic", "consensus_stance", "uncertain", "notes", "annotator"]


def fetch_top_level_comments(client, post_id: str, scanned: int) -> list[dict]:
    """The first `scanned` top-level comments of a post, oldest first."""
    found, seen, cursor = [], set(), None
    while len(found) < scanned:
        params = {"link_id": post_id, "limit": 100, "sort": "asc"}
        if cursor is not None:
            # One second back, as in reddit_archive.fetch_window; repeated IDs are dropped.
            params["after"] = cursor - 1
        page = reddit_archive._get(client, COMMENTS_URL, params)
        new = [comment for comment in page if comment.get("id") not in seen]
        seen.update(comment.get("id") for comment in new)
        found += [comment for comment in new if str(comment.get("parent_id", "")).startswith("t3_")]
        if len(page) < 100 or not new:
            break
        cursor = int(page[-1]["created_utc"])
    return found[:scanned]


def post_record(post: dict, topics: str) -> dict:
    return {**reddit_archive.post_row(post), "author": post.get("author", ""), "topics": topics}


def comment_record(comment: dict, post: dict, scanned: int, chosen: int) -> dict:
    created = comment.get("created_utc")
    body = comment.get("body", "")
    return {
        "subreddit": comment.get("subreddit", post.get("subreddit", "")), "id": comment.get("id", ""),
        "link_id": post.get("id", ""), "author": comment.get("author", ""), "body": body, "created_utc": created,
        "created_at_utc": datetime.fromtimestamp(created, tz=UTC).isoformat() if isinstance(created, (int, float)) else "",
        "score": comment.get("score", 0), "removed": int(body in REMOVED_TEXT), "day_weight": post.get("day_weight", 1.0),
        "scanned": scanned, "chosen": chosen, "post_topics": post.get("topics", ""),
    }


def choose(items: list, count: int, rng: np.random.Generator) -> list:
    if len(items) <= count:
        return list(items)
    return [items[index] for index in sorted(rng.choice(len(items), size=count, replace=False))]


def units(posts: pd.DataFrame, comments: pd.DataFrame) -> pd.DataFrame:
    """Posts and comments of the frame as units: post units read title, body and link; comment units the body."""
    introduced = pd.to_datetime(posts["created_at_utc"], utc=True, format="ISO8601")
    post_units = pd.DataFrame({
        "unit_id": "p_" + posts["id"].astype(str), "kind": "post", "author": posts["author"].fillna("").astype(str),
        "group": posts["subreddit"], "year": introduced.dt.year, "day": introduced.dt.strftime("%Y-%m-%d"),
        "introduced": posts["created_at_utc"], "title": posts["title"].fillna(""), "text": posts["text"].fillna(""),
        "link": posts["url"].fillna(""), "day_weight": posts["day_weight"].astype(float),
        "removed": (posts["removed_by_category"].fillna("").astype(str).ne("") | posts["text"].isin(REMOVED_TEXT)).astype(int),
        "topics": posts["topics"].fillna(""),
    })
    post_units["weight"] = post_units["day_weight"]
    when = pd.to_datetime(comments["created_at_utc"], utc=True, format="ISO8601")
    comment_units = pd.DataFrame({
        "unit_id": "c_" + comments["id"].astype(str), "kind": "comment", "author": comments["author"].fillna("").astype(str),
        "group": comments["subreddit"], "year": when.dt.year, "day": when.dt.strftime("%Y-%m-%d"),
        "introduced": comments["created_at_utc"], "title": "", "text": comments["body"].fillna(""), "link": "",
        "day_weight": comments["day_weight"].astype(float), "removed": comments["removed"].astype(int),
        "topics": comments["post_topics"].fillna(""),
    })
    comment_units["weight"] = comment_units["day_weight"] * comments["scanned"].astype(float) / comments["chosen"].astype(float)
    return pd.concat([post_units, comment_units], ignore_index=True)[UNIT_COLUMNS]


def analysis_sample(frame: pd.DataFrame, period: list[int], per_cell: int, seed: int) -> pd.DataFrame:
    """At most `per_cell` visible units per subreddit, year and kind, with weights back to the frame."""
    eligible = frame[frame["year"].between(*period) & frame["removed"].eq(0)]
    rng = np.random.default_rng(seed)
    parts = []
    for _, cell in eligible.groupby(["group", "year", "kind"], sort=True):
        chosen = cell if len(cell) <= per_cell else cell.sample(n=per_cell, random_state=rng)
        parts.append(chosen.assign(weight=chosen["weight"] * cell["weight"].sum() / chosen["weight"].sum()))
    return pd.concat(parts, ignore_index=True) if parts else eligible.iloc[0:0]


def panel_authors(sample: pd.DataFrame, min_years: int) -> pd.Index:
    """Authors (people, not bots or deleted accounts) with units in at least `min_years` different years."""
    people = sample[~sample["author"].str.lower().isin(NOT_AUTHORS)]
    years = people.groupby("author")["year"].nunique()
    return years.index[years >= min_years]


def pseudonym(authors: pd.Series, salt: str) -> pd.Series:
    """A stable pseudonym per author (12 hex characters of a salted SHA-256), so samples in git hold no usernames."""
    import hashlib

    return authors.fillna("").astype(str).map(lambda name: hashlib.sha256(f"{salt}:{name}".encode()).hexdigest()[:12] if name else "")


def validation_sample(sample: pd.DataFrame, periods: list[list[int]], per_period: int, seed: int) -> pd.DataFrame:
    """The same number of units drawn at random from every period, a third of each to `dev` (tuning) and the rest
    to `test`, in a random presentation order."""
    rng = np.random.default_rng(seed)
    parts = []
    for first, last in periods:
        part = sample[sample["year"].between(first, last)]
        chosen = part if len(part) <= per_period else part.sample(n=per_period, random_state=rng)
        split = np.where(rng.permutation(len(chosen)) < len(chosen) // 3, "dev", "test")
        parts.append(chosen.assign(period=f"{first}-{last}", inclusion=len(chosen) / max(len(part), 1), split=split))
    drawn = pd.concat(parts, ignore_index=True)
    return drawn.assign(order=rng.permutation(len(drawn)) + 1).sort_values("order", ignore_index=True)


def write_validation_workbook(path: Path, annotator: str, drawn: pd.DataFrame, codebook: Codebook) -> None:
    """Blind workbook for `consensus_topic` and `consensus_stance` of the H1 v2 validation sample."""
    workbook = Workbook()
    workbook.properties.description = json.dumps({"phase": "h1v2_validation", "annotator": annotator,
                                                  "codebook_version": codebook.version})
    notes = workbook.active
    notes.title = "Instrukcja"
    rules = {rule.code: rule.text for rule in codebook.rules}
    for row in [[f"Arkusz walidacyjny H1 v2, annotator {annotator}. Codebook v{codebook.version}."], [],
                ["consensus_topic: " + ", ".join(codebook.variable("consensus_topic").values)],
                ["consensus_stance: " + ", ".join(codebook.variable("consensus_stance").values)],
                ["R14: " + rules["R14"]], ["R15: " + rules["R15"]],
                ["Komentarz oceniamy w kontekście tematu posta (kolumna post_topics), nie otwierając linków."]]:
        notes.append(row)
    sheet = workbook.create_sheet("Reddit")
    header = ["unit_id", "rodzaj", "subreddit", "rok", "tytuł", "treść", "post_topics", *VALIDATION_LABELS]
    sheet.append(header)
    for unit in drawn.itertuples(index=False):
        sheet.append([unit.unit_id, unit.kind, unit.group, int(unit.year), str(unit.title)[:EXCEL_CELL_LIMIT],
                      str(unit.text)[:EXCEL_CELL_LIMIT], unit.topics, None, None, None, None, annotator])
    last = max(len(drawn) + 1, 2)
    for name, values in [("consensus_topic", codebook.variable("consensus_topic").values),
                         ("consensus_stance", codebook.variable("consensus_stance").values)]:
        letter = get_column_letter(header.index(name) + 1)
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, showErrorMessage=True)
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last}")
    uncertain = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
    sheet.add_data_validation(uncertain)
    letter = get_column_letter(header.index("uncertain") + 1)
    uncertain.add(f"{letter}2:{letter}{last}")
    for column, width in zip(range(1, len(header) + 1), [12, 9, 14, 6, 40, 80, 16, 16, 16, 10, 30, 10]):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if header[cell.column - 1] in VALIDATION_LABELS[:-1]:
                cell.fill = LABEL_FILL
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    sheet.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_validation_workbook(path: Path, codebook: Codebook) -> tuple[pd.DataFrame, list[str]]:
    """Labels of a filled validation workbook (rows without labels skipped) and the problems found."""
    rows = list(load_workbook(path, read_only=True, data_only=True)["Reddit"].iter_rows(values_only=True))
    header, labels, problems = list(rows[0]), [], []
    allowed = {name: set(codebook.variable(name).values) for name in ["consensus_topic", "consensus_stance"]}
    for number, values in enumerate(rows[1:], start=2):
        record = dict(zip(header, values))
        if not any(record.get(name) not in (None, "") for name in ["consensus_topic", "consensus_stance"]):
            continue
        label = {"unit_id": str(record["unit_id"]), "annotator": record["annotator"]}
        for name, values_allowed in allowed.items():
            value = str(record.get(name) or "").strip().lower()
            if value not in values_allowed:
                problems.append(f"wiersz {number} ({record['unit_id']}): {name} {value!r} spoza listy")
                value = ""
            label[name] = value
        flag, problem = _binary(record.get("uncertain"), required=False)
        if problem:
            problems.append(f"wiersz {number} ({record['unit_id']}): uncertain {problem}")
        label["uncertain"] = 0 if pd.isna(flag) else flag
        label["notes"] = str(record.get("notes") or "")
        labels.append(label)
    return pd.DataFrame(labels, columns=["unit_id", "annotator", "consensus_topic", "consensus_stance", "uncertain", "notes"]), problems
