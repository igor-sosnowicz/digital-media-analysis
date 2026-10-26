"""Error analysis of the instruments (slide 3): every error with a suggested cause, for people to confirm.

`error_table` lists one row per error of an instrument on gold units:
- detection errors of `cites_source` and `edit_about_sources` (false positive or negative);
- type errors of each source category, only where both gold and the instrument
  see a citation, so a missed citation is counted once, as a detection error.
`suggest_cause` proposes a cause from the closed list `CAUSES` from surface
features of the unit; it is a first guess. A person confirms or corrects it in
the review workbook (`annotations/sheets/errors_<split>.xlsx`, with the text),
and the confirmed causes go back to `annotations/errors/<split>_errors.csv`.
The Pareto table counts confirmed causes, or suggested ones while unconfirmed.
"""

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from digital_media_analysis.annotation_sheets import HEADER_FONT, LABEL_FILL, VERIFIED, unverified
from digital_media_analysis.codebook import NO_SOURCE
from digital_media_analysis.instruments import rules_v1
from digital_media_analysis.instruments.unit_text import OWN_MEDIA_HOSTS, link_host
from digital_media_analysis.validation import label_sets

CAUSES = [
    "podciąg", "brak kontekstu", "link niebędący źródłem", "skrót WP:", "znacznik edycji",
    "wyciek kontekstu komentarza", "brak w słowniku", "pominięte pole url", "inne",
]
ERROR_COLUMNS = [
    "unit_id", "platform", "split", "instrument", "variable", "error", "gold", "predicted", "suggested_cause", "cause", "notes",
]
SHORTCUT = re.compile(r"\[\[\s*(?:WP|WPT|WT|Pomoc|Help)\s*:", re.IGNORECASE)
EDIT_TAG = re.compile(r"źródła/przypisy|\bprzypis\w*|dodano przypis", re.IGNORECASE)
URL = re.compile(r"https?://", re.IGNORECASE)
FORM_HOSTS = {"forms.gle", "docs.google.com", "surveymonkey.com", "typeform.com"}
# v1 terms per codebook category, to tell substring matches from whole-word ones.
V1_TERMS: dict[str, set[str]] = {}
for _patterns, _mapping in [(rules_v1.REDDIT_REFERENCE_PATTERNS, rules_v1.REDDIT_CATEGORY_MAP),
                            (rules_v1.WIKIPEDIA_SOURCE_CATEGORIES, rules_v1.WIKIPEDIA_CATEGORY_MAP)]:
    for _category, _terms in _patterns.items():
        if (_code := _mapping.get(_category)) is not None:
            V1_TERMS.setdefault(_code, set()).update(term.casefold() for term in _terms)


def _substring_only(text: str, terms: set[str]) -> bool:
    lowered = text.casefold()
    hits = [term for term in terms if term in lowered]
    return bool(hits) and not any(re.search(rf"(?<!\w){re.escape(term)}(?!\w)", lowered) for term in hits)


def _whole_word(text: str, terms: set[str]) -> bool:
    lowered = text.casefold()
    return any(re.search(rf"(?<!\w){re.escape(term)}(?!\w)", lowered) for term in terms)


def _text(value) -> str:
    return "" if value is None or (not isinstance(value, str) and pd.isna(value)) else str(value)


def _unit_info(units: pd.DataFrame) -> dict[tuple[str, str], dict]:
    return {(row["platform"], str(row["unit_id"])): row for row in units.to_dict("records")}


def suggest_cause(platform: str, instrument: str, variable: str, error: str, text: str, link: str = "") -> str:
    """A first guess at the cause of one error, from the closed list `CAUSES`."""
    if instrument == "refs_v1":
        # It reads the article, not the summary: its errors are edits whose summary is silent about a
        # change of references, or declares work on them without one; none of the listed causes.
        return "inne"
    category = variable.split(":", 1)[1] if variable.startswith("source_types:") else None
    v1 = instrument == "rules_v1"
    if error == "FN":
        if platform == "reddit" and link and v1 and variable == "cites_source":
            return "pominięte pole url"
        return "brak w słowniku"
    if platform == "wikipedia":
        if SHORTCUT.search(text) and category in (None, "encyclopedia", "other"):
            return "skrót WP:"
        if EDIT_TAG.search(text) and category in (None, "unknown"):
            return "znacznik edycji"
    host = link_host(link) if link else ""
    if platform == "reddit" and (host in OWN_MEDIA_HOSTS or host in FORM_HOSTS):
        return "link niebędący źródłem"
    if platform == "wikipedia" and URL.search(text) and variable == "cites_source":
        return "link niebędący źródłem"
    if v1:
        terms = V1_TERMS.get(category, set()) if category else set().union(*V1_TERMS.values())
        if _substring_only(text, terms):
            return "podciąg"
        if platform == "wikipedia" and category and (URL.search(text) or SHORTCUT.search(text)) and _whole_word(text, terms):
            return "wyciek kontekstu komentarza"
    return "brak kontekstu"


def error_table(reference: pd.DataFrame, predictions: pd.DataFrame, units: pd.DataFrame, instrument: str,
                categories: list[str], split: str = "") -> pd.DataFrame:
    """Every error of one instrument on the gold units, with a suggested cause (columns `ERROR_COLUMNS`)."""
    gold = reference[["unit_id", "platform", "cites_source", "edit_about_sources", "source_types"]].astype({"unit_id": str})
    predicted = predictions[["unit_id", "platform", "cites_source", "edit_about_sources", "source_types"]].astype({"unit_id": str})
    frame = gold.merge(predicted, on=["unit_id", "platform"], suffixes=("", "_pred"))
    info = _unit_info(units)
    rows = []
    for unit in frame.itertuples(index=False):
        record = info.get((unit.platform, unit.unit_id), {})
        title, body, link = (_text(record.get(column)) for column in ["title", "text", "link"])
        content = f"{title}\n{body}".strip() if unit.platform == "reddit" else body
        found = []
        for variable in ["cites_source", "edit_about_sources"]:
            truth, guess = getattr(unit, variable), getattr(unit, f"{variable}_pred")
            if pd.notna(truth) and pd.notna(guess) and int(truth) != int(guess):
                found.append((variable, int(truth), int(guess)))
        if pd.notna(unit.source_types_pred) and unit.cites_source == 1 and unit.cites_source_pred == 1:
            truth_set = label_sets(pd.Series([unit.source_types])).iloc[0]
            guess_set = label_sets(pd.Series([unit.source_types_pred])).iloc[0]
            found += [(f"source_types:{code}", int(code in truth_set), int(code in guess_set))
                      for code in categories if code != NO_SOURCE and (code in truth_set) != (code in guess_set)]
        for variable, truth, guess in found:
            error = "FP" if guess and not truth else "FN"
            rows.append({
                "unit_id": unit.unit_id, "platform": unit.platform, "split": split, "instrument": instrument,
                "variable": variable, "error": error, "gold": truth, "predicted": guess,
                "suggested_cause": suggest_cause(unit.platform, instrument, variable, error, content, link),
                "cause": "", "notes": "",
            })
    return pd.DataFrame(rows, columns=ERROR_COLUMNS)


def write_review(path: Path, errors: pd.DataFrame, units: pd.DataFrame) -> None:
    """Review workbook: each error with the unit text, the suggested cause and an empty `cause` to confirm."""
    info = _unit_info(units)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Błędy"
    header = ["unit_id", "platform", "instrument", "variable", "error", "gold", "predicted", "tekst", "link",
              "suggested_cause", "cause", "notes"]
    sheet.append(header)
    for error in errors.itertuples(index=False):
        record = info.get((error.platform, str(error.unit_id)), {})
        content = "\n".join(_text(record.get(column)) for column in ["title", "text"]).strip()[:2000]
        sheet.append([error.unit_id, error.platform, error.instrument, error.variable, error.error, error.gold, error.predicted,
                      content, _text(record.get("link")), error.suggested_cause, error.cause or None, error.notes or None])
    last = max(sheet.max_row, 2)
    validation = DataValidation(type="list", formula1='"' + ",".join(CAUSES) + '"', allow_blank=True, showErrorMessage=True)
    sheet.add_data_validation(validation)
    validation.add(f"K2:K{last}")
    for row in range(2, last + 1):
        sheet.cell(row, 11).fill = LABEL_FILL
        sheet.cell(row, 12).fill = LABEL_FILL
    for index, width in enumerate([14, 10, 12, 24, 6, 6, 8, 80, 30, 24, 24, 30], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:L{last}"
    notes = workbook.create_sheet("Instrukcja")
    notes.append(["Analiza błędów: potwierdź albo popraw przyczynę każdego błędu."])
    notes.append(["Kolumna cause: wybierz z listy; puste = przyjmujemy suggested_cause jako niepotwierdzoną. Kolumna notes: uwagi."])
    notes.append(["Przyczyny:", ", ".join(CAUSES)])
    notes["A1"].font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_review(path: Path, errors: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """`errors` with `cause` and `notes` from the review workbook, and the problems found (unknown causes)."""
    rows = list(load_workbook(path, read_only=True, data_only=True)["Błędy"].iter_rows(values_only=True))
    header = list(rows[0])
    reviewed = pd.DataFrame([dict(zip(header, values)) for values in rows[1:]])
    problems = []
    if reviewed.empty:
        return errors, problems
    if VERIFIED in reviewed.columns:
        # An AI-suggested cause counts as confirmed only once a person verified the row.
        checked = [not unverified(record) for record in reviewed.to_dict("records")]
        reviewed.loc[[not flag for flag in checked], "cause"] = ""
    reviewed = reviewed.assign(unit_id=reviewed["unit_id"].astype(str)).fillna({"cause": "", "notes": ""})
    reviewed["cause"] = reviewed["cause"].astype(str).str.strip()
    for row in reviewed[reviewed["cause"].ne("") & ~reviewed["cause"].isin(CAUSES)].itertuples(index=False):
        problems.append(f"{row.unit_id} ({row.instrument}, {row.variable}): nieznana przyczyna {row.cause!r}")
    keys = ["unit_id", "platform", "instrument", "variable"]
    merged = errors.drop(columns=["cause", "notes"]).merge(reviewed[[*keys, "cause", "notes"]], on=keys, how="left")
    return merged.fillna({"cause": "", "notes": ""})[ERROR_COLUMNS], problems


def pareto(errors: pd.DataFrame) -> pd.DataFrame:
    """Errors per instrument and cause (confirmed, else suggested), descending, with the cumulative share."""
    cause = errors["cause"].where(errors["cause"].fillna("").ne(""), errors["suggested_cause"])
    counted = errors.assign(final_cause=cause, confirmed=errors["cause"].fillna("").ne(""))
    table = (counted.groupby(["instrument", "final_cause"])
             .agg(errors=("unit_id", "size"), confirmed=("confirmed", "sum"))
             .reset_index().rename(columns={"final_cause": "cause"}))
    table = table.sort_values(["instrument", "errors", "cause"], ascending=[True, False, True], ignore_index=True)
    totals = table.groupby("instrument")["errors"].transform("sum")
    return table.assign(share=table["errors"] / totals, cumulative_share=table.groupby("instrument")["errors"].cumsum() / totals)


def plot_pareto(table: pd.DataFrame, path=None):
    """Horizontal bars of errors per cause, one panel per instrument, labelled with count and cumulative share."""
    import matplotlib.pyplot as plt

    from digital_media_analysis.agreement import GRID, INK, INK_SECONDARY, SERIES, SURFACE

    instruments = list(dict.fromkeys(table["instrument"]))
    rows = max(table["instrument"].value_counts()) if len(table) else 1
    figure, axes = plt.subplots(1, max(len(instruments), 1), figsize=(5.6 * max(len(instruments), 1), 0.42 * rows + 1.4),
                                squeeze=False, facecolor=SURFACE)
    for axis, instrument in zip(axes[0], instruments):
        part = table[table["instrument"].eq(instrument)].reset_index(drop=True)
        # The same slots in every panel, filled from the top, so bars keep one thickness.
        positions = np.arange(rows)[::-1][: len(part)]
        axis.set_facecolor(SURFACE)
        axis.set_ylim(-0.6, rows - 0.4)
        axis.barh(positions, part["errors"], height=0.55, color=SERIES, zorder=2)
        top = part["errors"].max()
        for position, row in zip(positions, part.itertuples(index=False)):
            label = f"{row.errors}  ·  {row.cumulative_share:.0%} łącznie".replace("%", " %")
            axis.text(row.errors + top * 0.02, position, label, va="center", fontsize=8, color=INK)
        axis.set_xlim(0, top * 1.45)
        axis.set_yticks(positions)
        axis.set_yticklabels(part["cause"], fontsize=9, color=INK)
        axis.grid(axis="x", color=GRID, linewidth=0.8)
        axis.set_axisbelow(True)
        axis.tick_params(colors=INK_SECONDARY, length=0)
        for side in ["top", "right", "left"]:
            axis.spines[side].set_visible(False)
        axis.spines["bottom"].set_color(GRID)
        confirmed = int(part["confirmed"].sum())
        axis.set_title(f"{instrument}: {int(part['errors'].sum())} błędów (potwierdzone: {confirmed})",
                       loc="left", fontsize=11, color=INK)
        axis.set_xlabel("liczba błędów", fontsize=9, color=INK_SECONDARY)
    figure.tight_layout()
    if path is not None:
        figure.savefig(path, dpi=200, facecolor=SURFACE)
    return figure
