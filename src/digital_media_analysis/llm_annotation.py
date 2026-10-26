"""Automated annotation with two independent LLM annotators, rule checks and a human review queue (plan B).

Two annotators label the same rows from the codebook alone: A1 is Claude through
the Anthropic API, A2 a local open model from another family. Deterministic
rule checks (`rule_flags`) catch codebook violations. `combine` accepts a label
automatically only when both annotators agree, both are confident, no rule
objects and the label is not a rare one that people always check; every other
row, and a random control share of the accepted ones, goes to a person. Each
row keeps its review probability (`p_review`), so estimates from the reviewed
rows can be weighted back to all rows (two-phase sampling).

The prompts are built from codebook/codebook.yaml and are deterministic, so a
cached answer is reused only for the same model, prompt and unit. Answers are
cached in data/llm_cache/ (outside git); the text-free label files go to
annotations/llm/. Rules and procedure: docs/hipotezy.md (amendment of
26.10.2026) and docs/plan_b_automatyzacja.md.
"""

import hashlib
import json
import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from digital_media_analysis.codebook import NO_SOURCE, Codebook
from digital_media_analysis.errors import CAUSES
from digital_media_analysis.instruments.labels import LABEL_SEPARATOR
from digital_media_analysis.instruments.unit_text import OWN_MEDIA_HOSTS, link_host
from digital_media_analysis.paths import DATA_DIR

CACHE_DIR = DATA_DIR / "llm_cache"
CONFIDENCE = ["low", "medium", "high"]
CONTROL_SHARE = 0.15
SEED = 20261026

CAUSE_MEANINGS = {
    "podciąg": "słowo dopasowane wewnątrz innego słowa (np. „gus” w „Ferguson”)",
    "brak kontekstu": "nazwa lub kategoria pasuje, ale nie pełni funkcji źródła, np. instytucja jako temat (R2)",
    "link niebędący źródłem": "własne zdjęcie, ankieta, reklama, spam, profil (R3)",
    "skrót WP:": "skrót zasad lub narzędzi Wikipedii potraktowany jako źródło (R9)",
    "znacznik edycji": "znacznik „źródła/przypisy” potraktowany jako powołanie (R10)",
    "wyciek kontekstu komentarza": "kategoria wzięta ze słów z innej części opisu niż wzmianka o źródle",
    "brak w słowniku": "źródło nazwane lub podlinkowane, ale instrument go nie rozpoznał (np. domeny nie ma na liście)",
    "pominięte pole url": "instrument nie przeczytał linku posta (reguły v1)",
    "inne": "żadne z powyższych, także gdy to etykieta referencyjna (gold) wydaje się niezgodna z codebookiem",
}


@dataclass(frozen=True)
class Task:
    """What an annotator labels: the sheet columns it reads and the column it fills."""

    name: str
    inputs: tuple[str, ...]
    target: str
    multilabel: bool = False
    binary: str = ""
    # Labels a person always checks (rare classes that decide a validity gate).
    always_review: tuple[str, ...] = ()


TASKS = {
    "stance": Task("stance", ("subreddit", "tytuł", "treść", "link"), "consensus_stance", always_review=("challenges",)),
    "sources_reddit": Task("sources_reddit", ("subreddit", "tytuł", "treść", "link"), "source_types", multilabel=True),
    "sources_wikipedia": Task("sources_wikipedia", ("strona", "opis zmiany"), "source_types", multilabel=True,
                              binary="edit_about_sources"),
    "error_cause": Task("error_cause", ("platform", "instrument", "variable", "error", "gold", "predicted", "tekst", "link"),
                        "cause", always_review=("inne",)),
    "domain": Task("domain", ("domain", "links", "category"), "reviewed_category", always_review=("social_blog", "unknown")),
}


def values(task: Task, codebook: Codebook) -> list[str]:
    """The labels an annotator may give for a task."""
    if task.name == "stance":
        return list(codebook.variable("consensus_stance").values)
    if task.name == "error_cause":
        return list(CAUSES)
    codes = list(codebook.category_codes)
    if task.name == "domain":
        return [code for code in codes if code not in {NO_SOURCE, "personal"}]
    return codes


def _rules(codebook: Codebook, codes: Sequence[str]) -> list[str]:
    return [f"{rule.code}: {' '.join(rule.text.split())}" for rule in codebook.rules if rule.code in codes]


def system_prompt(task: Task, codebook: Codebook) -> str:
    """The annotation instructions for a task, built from the codebook only (no examples from the data)."""
    lines = [
        "Jesteś annotatorem w projekcie badawczym o powoływaniu się na źródła i zaufaniu do nauki w mediach społecznościowych.",
        "Etykietujesz jedną jednostkę naraz, ściśle według poniższego codebooka. Nie otwierasz linków: oceniasz tylko podany tekst i adres.",
        f"Konstrukt: {' '.join(codebook.construct_definition.split())}",
        "",
    ]
    if task.name == "stance":
        variable = codebook.variable("consensus_stance")
        lines += [f"Pytanie: {' '.join(variable.question.split())}", "Wartości:"]
        lines += [f"- {value}: {' '.join(meaning.split())}" for value, meaning in variable.values.items()]
        lines += ["", "Reguły:", *_rules(codebook, ["R6", "R11", "R12", "R14"])]
    elif task.name in {"sources_reddit", "sources_wikipedia", "domain"}:
        platform = "reddit" if task.name == "sources_reddit" else "wikipedia"
        if task.name != "domain":
            unit = next(unit for unit in codebook.units if unit.platform == platform)
            lines += [f"Jednostka: {unit.unit} ({' '.join(unit.text.split())}). {' '.join(unit.note.split())}"]
            lines += [f"Pytanie: {' '.join(codebook.variable('source_types').question.split())}"]
        else:
            lines += ["Zadanie: przypisz domenie internetowej kategorię źródła, czyli typ wydawcy lub kanału, który ją prowadzi."]
        lines += ["Kategorie:"]
        for category in codebook.categories:
            if category.code not in values(task, codebook):
                continue
            lines.append(f"- {category.code} ({category.name}): {' '.join(category.definition.split())}")
            lines += [f"    obejmuje: {' '.join(item.split())}" for item in category.include]
            lines += [f"    nie obejmuje: {' '.join(item.split())}" for item in category.exclude]
        if task.binary:
            variable = codebook.variable(task.binary)
            lines += ["", f"Zmienna {task.binary}: {' '.join(variable.question.split())}"]
            lines += [f"- {value}: {' '.join(meaning.split())}" for value, meaning in variable.values.items()]
        lines += ["", "Reguły:", *_rules(codebook, [rule.code for rule in codebook.rules if rule.code not in {"R14", "R15"}])]
    elif task.name == "error_cause":
        lines += ["Zadanie: instrument automatyczny pomylił się na jednostce. Wskaż najbardziej prawdopodobną przyczynę błędu.",
                  "FP: instrument przypisał coś, czego nie ma w etykiecie referencyjnej (gold). FN: pominął coś, co gold ma.",
                  "Przyczyny:"]
        lines += [f"- {cause}: {meaning}" for cause, meaning in CAUSE_MEANINGS.items()]
        lines += ["", "Kategorie źródeł z codebooka: " + ", ".join(codebook.category_codes),
                  "", "Reguły:", *_rules(codebook, [rule.code for rule in codebook.rules if rule.code not in {"R14", "R15"}])]
    lines += ["", "Odpowiedz wyłącznie obiektem JSON zgodnym ze schematem: etykieta, pewność (low, medium, high) i jedno zdanie "
              "uzasadnienia po polsku, najlepiej z numerem reguły. Przy wątpliwości wybierz najlepszą etykietę i daj pewność low."]
    return "\n".join(lines)


def unit_prompt(task: Task, record: dict) -> str:
    """The unit as the annotator sees it: the task's input columns, empty ones left out."""
    parts = []
    for column in task.inputs:
        value = record.get(column)
        if value is None or (isinstance(value, float) and np.isnan(value)) or str(value).strip() in {"", "nan", "None"}:
            continue
        parts.append(f"{column}: {str(value).strip()}")
    return "\n".join(parts) or "(pusta jednostka)"


def schema(task: Task, codebook: Codebook) -> dict:
    """JSON schema of an answer."""
    label = ({"type": "array", "items": {"type": "string", "enum": values(task, codebook)}, "minItems": 1}
             if task.multilabel else {"type": "string", "enum": values(task, codebook)})
    properties = {"label": label, "confidence": {"type": "string", "enum": CONFIDENCE}, "reason": {"type": "string"}}
    if task.binary:
        properties[task.binary] = {"type": "integer", "enum": [0, 1]}
    return {"type": "object", "properties": properties, "required": list(properties), "additionalProperties": False}


def normalize(task: Task, answer: dict | None, codebook: Codebook) -> dict:
    """A validated answer as {label, binary, confidence, reason}; an unusable one has label None."""
    empty = {"label": None, "binary": None, "confidence": "low", "reason": "brak poprawnej odpowiedzi"}
    if not isinstance(answer, dict):
        return empty
    allowed = set(values(task, codebook))
    raw = answer.get("label")
    if task.multilabel:
        codes = sorted({str(code).strip().lower() for code in (raw if isinstance(raw, list) else str(raw or "").split(LABEL_SEPARATOR))
                        if str(code).strip()})
        if not codes or any(code not in allowed for code in codes):
            return empty
        if NO_SOURCE in codes and len(codes) > 1:
            codes = [code for code in codes if code != NO_SOURCE]
        label = LABEL_SEPARATOR.join(codes)
    else:
        label = str(raw or "").strip()
        if label not in allowed:
            return empty
    binary = answer.get(task.binary) if task.binary else None
    if task.binary and binary not in (0, 1):
        return empty
    confidence = answer.get("confidence") if answer.get("confidence") in CONFIDENCE else "low"
    return {"label": label, "binary": binary, "confidence": confidence, "reason": str(answer.get("reason") or "")[:300]}


def _digest(*parts: str) -> str:
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()


@dataclass
class Annotator:
    """An LLM annotator with a JSON-lines answer cache; `ask(system, prompt, schema) -> (answer dict | None, model id)`."""

    name: str
    model: str
    ask: Callable[[str, str, dict], tuple[dict | None, str]]
    cache_dir: Path = CACHE_DIR
    _cache: dict = field(default_factory=dict, init=False)

    def _cache_path(self) -> Path:
        return self.cache_dir / f"{self.name}.jsonl"

    def _load(self) -> None:
        if self._cache or not self._cache_path().exists():
            return
        for line in self._cache_path().read_text(encoding="utf-8").splitlines():
            if line.strip():
                entry = json.loads(line)
                self._cache[entry["key"]] = entry

    def annotate(self, task: Task, records: Sequence[dict], key_columns: Sequence[str], codebook: Codebook,
                 progress: Callable[[int, int], None] | None = None) -> pd.DataFrame:
        """One row per record: its key columns, the label (and binary variable), confidence, reason, model and prompt hash."""
        self._load()
        system, answer_schema = system_prompt(task, codebook), schema(task, codebook)
        system_sha = _digest(system)[:12]
        rows = []
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        with self._cache_path().open("a", encoding="utf-8") as cache:
            for index, record in enumerate(records):
                prompt = unit_prompt(task, record)
                key = _digest(self.model, system, prompt)
                if key not in self._cache:
                    answer, served_by = self.ask(system, prompt, answer_schema)
                    entry = {"key": key, "model": self.model, "served_by": served_by, "answer": answer}
                    cache.write(json.dumps(entry, ensure_ascii=False) + "\n")
                    cache.flush()
                    self._cache[key] = entry
                entry = self._cache[key]
                result = normalize(task, entry["answer"], codebook)
                rows.append({**{column: str(record[column]) for column in key_columns},
                             "label": result["label"], "binary": result["binary"], "confidence": result["confidence"],
                             "reason": result["reason"], "model": entry.get("served_by") or self.model,
                             "system_sha": system_sha})
                if progress:
                    progress(index + 1, len(records))
        return pd.DataFrame(rows)


def claude_annotator(model: str = "claude-opus-5-5", effort: str = "low", cache_dir: Path = CACHE_DIR) -> Annotator:
    """A1: Claude through the Anthropic API (credentials from the environment), structured JSON output."""
    import anthropic

    client = anthropic.Anthropic()

    def ask(system: str, prompt: str, answer_schema: dict) -> tuple[dict | None, str]:
        response = client.beta.messages.create(
            model=model,
            max_tokens=4096,
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            system=[{"type": "text", "text": system, "cache_control": {"type": "ephemeral"}}],
            output_config={"effort": effort, "format": {"type": "json_schema", "schema": answer_schema}},
            messages=[{"role": "user", "content": prompt}],
        )
        if response.stop_reason == "refusal":
            return None, response.model
        text = next((block.text for block in response.content if block.type == "text"), "")
        try:
            return json.loads(text), response.model
        except json.JSONDecodeError:
            return None, response.model

    return Annotator(f"A1_{model}", model, ask, cache_dir)


def local_annotator(model: str = "Qwen/Qwen2.5-3B-Instruct", cache_dir: Path = CACHE_DIR) -> Annotator:
    """A2: a local open instruction model (transformers, `ml` group); the JSON is parsed from its reply."""
    from transformers import pipeline

    generate = pipeline("text-generation", model=model)

    def ask(system: str, prompt: str, answer_schema: dict) -> tuple[dict | None, str]:
        messages = [{"role": "system", "content": system + "\nSchemat JSON: " + json.dumps(answer_schema, ensure_ascii=False)},
                    {"role": "user", "content": prompt}]
        output = generate(messages, max_new_tokens=200, do_sample=False)[0]["generated_text"][-1]["content"]
        found = re.search(r"\{.*\}", output, re.DOTALL)
        try:
            return (json.loads(found.group(0)) if found else None), model
        except json.JSONDecodeError:
            return None, model

    slug = model.replace("/", "_")
    return Annotator(f"A2_{slug}", model, ask, cache_dir)


REVERT = re.compile(r"Anulowanie wersji|Wycofano (?:edycj|ostatni)|^Bot:|komunikat bota", re.IGNORECASE)
REFERENCE_WORK = re.compile(r"źródła/przypisy|dodan\w* przypis|przypis\w*|źródło\b", re.IGNORECASE)


def rule_flags(task: Task, record: dict, label: str | None, binary: int | None) -> list[str]:
    """Codebook rules the label breaks, checked without any model (empty list when none)."""
    if label is None:
        return []
    flags = []
    codes = set(label.split(LABEL_SEPARATOR)) if task.multilabel else {label}
    if task.multilabel and NO_SOURCE in codes and len(codes) > 1:
        flags.append("R13")
    if task.name == "sources_reddit":
        text = " ".join(str(record.get(column) or "") for column in ("tytuł", "treść")).strip()
        host = link_host(str(record.get("link") or "")) if str(record.get("link") or "") not in {"", "nan", "None"} else ""
        if host in OWN_MEDIA_HOSTS and codes != {NO_SOURCE} and "http" not in text:
            flags.append("R3")
    if task.name == "sources_wikipedia":
        text = str(record.get("opis zmiany") or "")
        if REVERT.search(text) and codes != {NO_SOURCE}:
            flags.append("R9")
        if "źródła/przypisy" in text and binary == 0:
            flags.append("R10")
    return flags


def combine(first: pd.DataFrame, second: pd.DataFrame, records: pd.DataFrame, task: Task, key_columns: Sequence[str],
            control_share: float = CONTROL_SHARE, seed: int = SEED) -> pd.DataFrame:
    """Per row: both answers, rule flags, the suggestion, whether it is accepted, why it goes to a person, and p_review."""
    keys = list(key_columns)
    # Only the rows of `records` (the sheet): answers for other rows are ignored.
    both = records.astype({column: str for column in keys}).merge(
        first.merge(second, on=keys, suffixes=("_a1", "_a2"), how="outer"), on=keys, how="left")
    rng = np.random.default_rng(seed)
    rows = []
    for record in both.to_dict("records"):
        label1, label2 = record.get("label_a1"), record.get("label_a2")
        label1 = None if pd.isna(label1) else label1
        label2 = None if pd.isna(label2) else label2
        binary1 = None if pd.isna(record.get("binary_a1")) else int(record["binary_a1"])
        binary2 = None if pd.isna(record.get("binary_a2")) else int(record["binary_a2"])
        flags = sorted(set(rule_flags(task, record, label1, binary1)) | set(rule_flags(task, record, label2, binary2)))
        agree = label1 is not None and label1 == label2 and binary1 == binary2
        confident = record.get("confidence_a1") != "low" and record.get("confidence_a2") != "low"
        suggestion = label1 if label1 is not None else label2
        codes = set(str(suggestion).split(LABEL_SEPARATOR)) if task.multilabel else {suggestion}
        rare = bool(codes & set(task.always_review))
        if not agree:
            reason = "spór A1 i A2" if label1 is not None and label2 is not None else "brak odpowiedzi"
        elif flags:
            reason = "reguła " + ", ".join(flags)
        elif not confident:
            reason = "niska pewność"
        elif rare:
            reason = "rzadka etykieta"
        elif task.name == "domain" and suggestion != record.get("category"):
            # The domain list was classified by an AI assistant before; a change of category needs a person.
            reason = "zmiana kategorii"
        else:
            reason = ""
        accepted = not reason
        control = accepted and rng.random() < control_share
        rows.append({
            **{column: record[column] for column in keys},
            "label_a1": label1, "label_a2": label2, "binary_a1": binary1, "binary_a2": binary2,
            "confidence_a1": record.get("confidence_a1"), "confidence_a2": record.get("confidence_a2"),
            "flags": ";".join(flags), "agree": agree, "suggestion": suggestion,
            "suggestion_binary": binary1 if binary1 is not None else binary2,
            "decided_by": "human" if (not accepted or control) else "auto",
            "queue_reason": reason or ("kontrola" if control else ""),
            "p_review": 1.0 if not accepted else control_share,
            "reason": f"A1 ({label1}, {record.get('confidence_a1')}): {record.get('reason_a1') or ''} | "
                      f"A2 ({label2}, {record.get('confidence_a2')}): {record.get('reason_a2') or ''}"[:500],
        })
    return pd.DataFrame(rows)


def read_rows(path: Path, sheet: str | None = None) -> pd.DataFrame:
    """The rows of a workbook sheet (header in the first row) or of a CSV file, as text."""
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    from openpyxl import load_workbook

    rows = list(load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    frame = pd.DataFrame(rows[1:], columns=rows[0])
    return frame[frame.iloc[:, 0].notna()].reset_index(drop=True)


def write_queue(template: Path, output: Path, sheet: str | None, queue: pd.DataFrame, task: Task, key_columns: Sequence[str],
                source: str) -> None:
    """The template with suggestions prefilled, the AI columns, and decided_by, queue_reason and p_review per row.

    Rows the queue accepted (`decided_by = auto`) count without a check; the rest wait for `verified = 1`.
    """
    keys = list(key_columns)
    suggestions = queue[keys].assign(**{f"ai_{task.target}": queue["suggestion"]})
    targets = [task.target]
    if task.binary:
        suggestions[f"ai_{task.binary}"] = queue["suggestion_binary"].astype("Int64")
        targets.append(task.binary)
    suggestions = suggestions.assign(ai_uncertain=queue["decided_by"].eq("human").astype(int), ai_reason=queue["reason"])
    extra = queue[keys + ["decided_by", "queue_reason", "p_review"]]
    output.parent.mkdir(parents=True, exist_ok=True)
    if template.suffix.lower() == ".csv":
        frame = read_rows(template).astype({column: str for column in keys})
        merged = frame.merge(suggestions.astype({column: str for column in keys}), on=keys, how="left")
        for target in targets:
            if target in merged:
                empty = merged[target].fillna("").astype(str).str.strip().eq("")
                merged.loc[empty, target] = merged.loc[empty, f"ai_{target}"]
        merged = merged.merge(extra.astype({column: str for column in keys}), on=keys, how="left").assign(verified="")
        merged.to_csv(output, index=False, encoding="utf-8-sig")
        return
    import shutil

    from openpyxl import load_workbook

    from digital_media_analysis import ai_suggestions

    shutil.copyfile(template, output)
    ai_suggestions.prefill(output, sheet, keys, suggestions, targets, source)
    workbook = load_workbook(output)
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1]]
    by_key = {tuple(str(value) for value in row[keys]): row for _, row in extra.iterrows()}
    first = len(header) + 1
    for offset, name in enumerate(["decided_by", "queue_reason", "p_review"]):
        worksheet.cell(1, first + offset, name)
    for row in range(2, worksheet.max_row + 1):
        key = tuple(str(worksheet.cell(row, header.index(column) + 1).value) for column in keys)
        if key in by_key:
            for offset, name in enumerate(["decided_by", "queue_reason", "p_review"]):
                worksheet.cell(row, first + offset, by_key[key][name])
    workbook.save(output)
