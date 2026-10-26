"""Write blind annotation workbooks and read the filled-in labels back.

One workbook per annotator holds an instruction sheet and one sheet per platform.
Instrument output is never included, so annotators code blind. Label columns use
the names from `annotation_columns` in the codebook. `cites_source` is not in
the sheet: reading derives it from `source_types` (codebook rule R13).
"""

import json
from collections.abc import Sequence
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from digital_media_analysis.codebook import NO_SOURCE, Codebook
from digital_media_analysis.instruments.labels import LABEL_SEPARATOR

SHEET_NAMES = {"reddit": "Reddit", "wikipedia": "Wikipedia"}
# (unit column, header, width) for the read-only part of each sheet.
CONTENT_COLUMNS = {
    "reddit": [("unit_id", "unit_id", 12), ("group", "subreddit", 16), ("title", "tytuł", 40), ("text", "treść", 90), ("link", "link", 30)],
    "wikipedia": [("unit_id", "unit_id", 28), ("group", "strona", 22), ("text", "opis zmiany", 90)],
}
LABEL_COLUMNS = {
    "reddit": ["source_types", "uncertain", "notes", "annotator"],
    "wikipedia": ["source_types", "edit_about_sources", "uncertain", "notes", "annotator"],
}
BINARY_COLUMNS = ["edit_about_sources", "uncertain"]
LABEL_FILE_COLUMNS = [
    "unit_id", "platform", "annotator", "codebook_version",
    "cites_source", "source_types", "edit_about_sources", "uncertain", "notes",
]
EXCEL_CELL_LIMIT = 32_000
HEADER_FONT = Font(bold=True)
LABEL_FILL = PatternFill("solid", fgColor="FFF2CC")
# Column a person sets to 1 after checking an AI suggestion (ai_suggestions.prefill).
VERIFIED = "verified"
# auto: accepted by the two LLM annotators without review; human: in the review queue.
DECIDED_BY = "decided_by"


def write_workbook(
    path: Path,
    annotator: str,
    phase: str,
    units_by_platform: dict[str, pd.DataFrame],
    codebook: Codebook,
    context_columns: Sequence[str] = (),
    extra_rules: Sequence[str] = (),
) -> None:
    """Write one annotator's workbook; units must already be in presentation order.

    `context_columns` are unit columns shown read-only to the right of the label
    columns (an arbiter sees the annotators' labels there); `extra_rules` are
    added to the working rules on the instruction sheet.
    """
    workbook = Workbook()
    workbook.properties.title = f"{phase}_{annotator}"
    workbook.properties.description = json.dumps(
        {"phase": phase, "annotator": annotator, "codebook_version": codebook.version}
    )
    _write_instructions(workbook.active, annotator, phase, codebook, extra_rules)
    for platform, units in units_by_platform.items():
        sheet = workbook.create_sheet(SHEET_NAMES[platform])
        _write_platform_sheet(sheet, platform, units, annotator, codebook, context_columns)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_instructions(sheet, annotator: str, phase: str, codebook: Codebook, extra_rules: Sequence[str] = ()) -> None:
    sheet.title = "Instrukcja"
    rules = [
        "Koduj samodzielnie. Nie porównuj etykiet z innymi przed zakończeniem całego arkusza.",
        "Opieraj się tylko na codebooku i tekście w wierszu. Nie otwieraj linków.",
        "Wypełnij żółte kolumny w każdym wierszu. Zmienne binarne: 0 lub 1 (lista rozwijana).",
        f"source_types: kody z tabeli poniżej rozdzielone znakiem „{LABEL_SEPARATOR}”; {NO_SOURCE}, gdy nie ma powołania (sam, bez innych kodów; R13).",
        "W razie wątpliwości wybierz najlepszą etykietę, ustaw uncertain = 1 i opisz problem w notes (R12).",
        "Nie zmieniaj kolumn unit_id ani kolejności wierszy.",
        *extra_rules,
    ]
    rows = [
        [f"Arkusz annotacji: {phase}, annotator {annotator}"],
        [f"Codebook v{codebook.version} z {codebook.date.isoformat()}. Pełna instrukcja: codebook/codebook.md"],
        [],
        ["Zasady pracy"],
        *([f"{number}. {rule}"] for number, rule in enumerate(rules, start=1)),
        [],
        ["Kod", "Kategoria", "Definicja"],
        *([category.code, category.name, category.definition] for category in codebook.categories),
        [],
        ["Kolumna", "Opis"],
        *([column.name, column.description] for column in codebook.annotation_columns),
    ]
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column == 1 and cell.value in {"Zasady pracy", "Kod", "Kolumna"}:
                for header_cell in row:
                    header_cell.font = HEADER_FONT
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 45
    sheet.column_dimensions["C"].width = 90


def _write_platform_sheet(
    sheet, platform: str, units: pd.DataFrame, annotator: str, codebook: Codebook, context_columns: Sequence[str] = ()
) -> None:
    content = CONTENT_COLUMNS[platform]
    labels = LABEL_COLUMNS[platform]
    sheet.append([header for _, header, _ in content] + labels + list(context_columns))
    _append_rows(sheet, platform, units, annotator, context_columns, start_row=2)
    for index, (_, _, width) in enumerate(content, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offset, name in enumerate(labels, start=len(content) + 1):
        sheet.column_dimensions[get_column_letter(offset)].width = 30 if name in {"source_types", "notes"} else 14
    for offset in range(len(content) + len(labels) + 1, len(content) + len(labels) + len(context_columns) + 1):
        sheet.column_dimensions[get_column_letter(offset)].width = 24
    _format_label_cells(sheet, platform, 2, len(units) + 1, codebook)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    sheet.freeze_panes = "B2"


def _append_rows(
    sheet, platform: str, units: pd.DataFrame, annotator: str, context_columns: Sequence[str] = (), start_row: int = 2
) -> None:
    """Write unit rows from `start_row` on, with empty labels, the annotator and any context columns."""
    content = CONTENT_COLUMNS[platform]
    labels = LABEL_COLUMNS[platform]
    for row_number, unit in enumerate(units.to_dict("records"), start=start_row):
        values = [str(unit[column])[:EXCEL_CELL_LIMIT] for column, _, _ in content]
        values += [None] * (len(labels) - 1) + [annotator]
        values += [_shown(unit[column]) for column in context_columns]
        for column_number, value in enumerate(values, start=1):
            sheet.cell(row_number, column_number, value)


def _shown(value) -> str | None:
    return None if value is None or (not isinstance(value, str) and pd.isna(value)) else str(value)


def _format_label_cells(sheet, platform: str, first_row: int, last_row: int, codebook: Codebook) -> None:
    """Fill, wrapping and data validation (0/1 lists, code prompt) of the label cells in rows first_row..last_row."""
    content = CONTENT_COLUMNS[platform]
    if last_row >= first_row:
        for offset, name in enumerate(LABEL_COLUMNS[platform], start=len(content) + 1):
            letter = get_column_letter(offset)
            for row in range(first_row, last_row + 1):
                sheet.cell(row, offset).fill = LABEL_FILL
            if name in BINARY_COLUMNS:
                validation = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True)
            elif name == "source_types":
                validation = DataValidation(allow_blank=True, showInputMessage=True)
                validation.promptTitle = "Kody kategorii"
                validation.prompt = LABEL_SEPARATOR.join(codebook.category_codes)
            else:
                continue
            sheet.add_data_validation(validation)
            validation.add(f"{letter}{first_row}:{letter}{last_row}")
        for row in sheet.iter_rows(min_row=first_row, max_row=last_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(last_row, 2)}"


def append_units(
    path: Path, units_by_platform: dict[str, pd.DataFrame], codebook: Codebook, counts: dict[str, int | None] | None = None
) -> dict[str, int]:
    """Add the units a workbook lacks below its last unit row; the rows already there, and their labels, stay.

    `units_by_platform` are the candidate units in presentation order; those
    already in the sheet are skipped and `counts` caps how many are added per
    platform. A platform sheet the workbook does not have yet is created.
    Returns the number of rows added per platform.
    """
    lock = path.with_name(f"~${path.name}")
    if lock.exists():
        raise RuntimeError(f"{path.name} is open in Excel ({lock.name} exists); close it and try again.")
    workbook = load_workbook(path)
    annotator = _metadata(workbook).get("annotator") or path.stem.split("_", 1)[-1]
    added = {}
    for platform, units in units_by_platform.items():
        count = (counts or {}).get(platform)
        sheet_name = SHEET_NAMES[platform]
        if sheet_name not in workbook.sheetnames:
            chosen = units if count is None else units.head(count)
            _write_platform_sheet(workbook.create_sheet(sheet_name), platform, chosen, annotator, codebook)
            added[platform] = len(chosen)
            continue
        sheet = workbook[sheet_name]
        header = [cell.value for cell in sheet[1]]
        expected = [name for _, name, _ in CONTENT_COLUMNS[platform]] + LABEL_COLUMNS[platform]
        if header[: len(expected)] != expected:
            raise ValueError(f"{path.name}, {sheet_name}: headers {header} differ from the expected {expected}.")
        ids = [row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)]
        present = {str(unit_id) for unit_id in ids if unit_id is not None}
        # The last row holding a unit; formatting Excel added below it does not count.
        last_row = 1 + max((number for number, unit_id in enumerate(ids, start=1) if unit_id is not None), default=0)
        missing = units[~units["unit_id"].astype(str).isin(present)]
        chosen = missing if count is None else missing.head(count)
        _append_rows(sheet, platform, chosen, annotator, start_row=last_row + 1)
        _format_label_cells(sheet, platform, last_row + 1, last_row + len(chosen), codebook)
        added[platform] = len(chosen)
    workbook.save(path)
    return added


def units_after(units: pd.DataFrame, coded: list[str], count: int | None = None) -> pd.DataFrame:
    """The next `count` units (all when None) after the last coded one, in presentation order.

    `units` are in presentation order; `coded` are the unit IDs another annotator
    labelled (`coded_unit_ids`). Coded units never reappear, and units before the
    last coded one are left to that annotator even if they skipped some.
    """
    ids = units["unit_id"].astype(str).tolist()
    if unknown := sorted(set(coded) - set(ids)):
        raise ValueError(f"{len(unknown)} coded units are not in this phase of the sample, e.g. {unknown[:3]}.")
    start = max((ids.index(unit_id) for unit_id in coded), default=-1) + 1
    rest = units.iloc[start:]
    rest = rest[~rest["unit_id"].astype(str).isin(set(coded))]
    return rest if count is None else rest.head(count)


def coded_unit_ids(path: Path) -> dict[str, list[str]]:
    """Unit IDs of the rows that carry any label, per platform, in sheet order."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    coded = {}
    for platform, sheet_name in SHEET_NAMES.items():
        if sheet_name not in workbook.sheetnames:
            continue
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        header = list(rows[0])
        records = (dict(zip(header, values)) for values in rows[1:])
        coded[platform] = [str(record["unit_id"]) for record in records if record["unit_id"] is not None and _has_labels(record)]
    return coded


def workbook_metadata(path: Path) -> dict:
    """Phase, annotator and codebook version stored by `write_workbook` ({} if absent)."""
    return _metadata(load_workbook(path, read_only=True))


def read_workbook(path: Path, codebook: Codebook, partial: bool = False) -> tuple[pd.DataFrame, list[str]]:
    """Return the labels of a filled workbook and a list of problems found.

    Rows with problems are still returned, with the offending value set to missing,
    so the problems can be fixed in the workbook and the file read again. With
    `partial`, rows without any label are skipped instead of reported, for a
    phase the annotators split between them; filled rows are checked as usual.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    metadata = _metadata(workbook)
    known_codes = set(codebook.category_codes)
    labeled, problems, unchecked = [], [], []
    for platform, sheet_name in SHEET_NAMES.items():
        if sheet_name not in workbook.sheetnames:
            continue
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        header, body = list(rows[0]), rows[1:]
        expected = [name for _, name, _ in CONTENT_COLUMNS[platform]] + LABEL_COLUMNS[platform]
        if header[: len(expected)] != expected:
            problems.append(f"{sheet_name}: nagłówki {header} różnią się od oczekiwanych {expected}.")
            continue
        for excel_row, values in enumerate(body, start=2):
            record = dict(zip(header, values))
            if record["unit_id"] is None:
                continue
            if unverified(record):
                # A prefilled label nobody has checked is the AI's, not the person's.
                unchecked.append(f"{sheet_name}: {record['unit_id']}")
                continue
            if partial and not _has_labels(record):
                continue
            where = f"{sheet_name}, wiersz {excel_row} ({record['unit_id']})"
            label = {
                "unit_id": str(record["unit_id"]),
                "platform": platform,
                "annotator": record["annotator"],
                "codebook_version": metadata.get("codebook_version", ""),
            }
            for name in BINARY_COLUMNS:
                if name not in record:
                    label[name] = pd.NA
                    continue
                value, problem = _binary(record[name], required=name != "uncertain")
                label[name] = value
                if problem:
                    problems.append(f"{where}: {name} {problem}")
            label["uncertain"] = 0 if pd.isna(label["uncertain"]) else label["uncertain"]

            raw_codes = str(record["source_types"] or "").replace(",", LABEL_SEPARATOR)
            codes = [code.strip().lower() for code in raw_codes.split(LABEL_SEPARATOR)]
            codes = sorted({code for code in codes if code})
            if unknown := [code for code in codes if code not in known_codes]:
                problems.append(f"{where}: nieznane kody {unknown}")
                codes = [code for code in codes if code in known_codes]
            label["source_types"] = LABEL_SEPARATOR.join(codes)
            # cites_source follows from source_types (R13); missing when the label is unusable.
            if not codes:
                if not unknown:
                    problems.append(f"{where}: source_types brak wartości (wpisz {NO_SOURCE}, gdy nie ma powołania)")
                label["cites_source"] = pd.NA
            elif NO_SOURCE in codes and len(codes) > 1:
                problems.append(f"{where}: {NO_SOURCE} razem z innymi kodami ({label['source_types']})")
                label["cites_source"] = pd.NA
            else:
                label["cites_source"] = int(codes != [NO_SOURCE])
            label["notes"] = "" if record["notes"] is None else str(record["notes"])
            labeled.append(label)
    if unchecked and not partial:
        problems.append(f"{len(unchecked)} wierszy z podpowiedzią AI bez verified = 1 (sprawdź etykietę i wpisz 1), "
                        f"np. {unchecked[:3]}")
    return pd.DataFrame(labeled, columns=LABEL_FILE_COLUMNS), problems


def unverified(record: dict) -> bool:
    """Whether a sheet row carries an AI suggestion (columns `ai_*`) the person has not checked (verified = 1) yet.

    A row the review queue accepted automatically (`decided_by = auto`, llm_annotation.combine) needs no check.
    """
    if str(record.get(DECIDED_BY) or "").strip() == "auto":
        return False
    suggested = any(isinstance(name, str) and name.startswith("ai_") and value not in (None, "") for name, value in record.items())
    if not suggested:
        return False
    value, _ = _binary(record.get(VERIFIED), required=False)
    return pd.isna(value) or value != 1


def _has_labels(record: dict) -> bool:
    """Whether the annotator filled any label cell of a sheet row (the annotator column does not count)."""
    return any(
        value is not None and str(value).strip() != ""
        for name, value in record.items()
        if name in LABEL_FILE_COLUMNS and name not in {"unit_id", "annotator"}
    )


def _metadata(workbook) -> dict:
    """Phase, annotator and codebook version stored by `write_workbook` ({} if absent)."""
    try:
        return json.loads(workbook.properties.description or "{}")
    except json.JSONDecodeError:
        return {}


def _binary(value, required: bool) -> tuple[object, str]:
    """Parse a 0/1 cell; return (value or NA, problem description or '')."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return pd.NA, "brak wartości" if required else ""
    try:
        number = float(str(value).strip().replace(",", "."))
    except ValueError:
        return pd.NA, f"ma wartość {value!r} zamiast 0 lub 1"
    if number not in (0, 1):
        return pd.NA, f"ma wartość {value!r} zamiast 0 lub 1"
    return int(number), ""
