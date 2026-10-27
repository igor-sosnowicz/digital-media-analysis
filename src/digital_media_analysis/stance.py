"""Stance pass: blind workbooks for `consensus_stance` on Reddit posts, their import, reference and agreement.

The stance of a post toward the scientific consensus (codebook variable
`consensus_stance`, rule R14) is coded in a pass of its own, so the source
workbooks of codebook 2.0 stay as they are. A stance workbook,
`annotations/sheets/stance_<phase>_<annotator>.xlsx`, holds the Reddit posts and
the columns consensus_stance, uncertain, notes and annotator. Imported labels go
to `annotations/labels/<phase>/stance_<annotator>.csv`, without post text.

A workbook may also carry AI suggestions (`add_ai_suggestions`): consensus_stance
is prefilled, the suggestion itself is kept in ai_stance, and a row counts as
coded only after the person checked it and set verified = 1.
"""

import json
import re
from collections.abc import Sequence
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from digital_media_analysis.agreement import cohen_kappa, krippendorff_alpha
from digital_media_analysis.annotation_sheets import (
    CONTENT_COLUMNS,
    EXCEL_CELL_LIMIT,
    HEADER_FONT,
    LABEL_FILL,
    _binary,
    _metadata,
)
from digital_media_analysis.annotation_sheets import unverified as is_unverified  # `unverified` names a list below
from digital_media_analysis.annotation_sheets import verified_by as sheet_verified_by
from digital_media_analysis.codebook import Codebook
from digital_media_analysis.paths import LABELS_DIR, SHEETS_DIR

STANCE = "consensus_stance"
SHEET_NAME = "Reddit"
LABEL_COLUMNS = [STANCE, "uncertain", "notes", "annotator"]
FILE_COLUMNS = ["unit_id", "platform", "annotator", "codebook_version", STANCE, "uncertain", "notes", "ai_stance",
                "decided_by", "p_review", "verified_by"]
# Columns after LABEL_COLUMNS in a workbook with AI suggestions; the person fills only `verified`.
AI_COLUMNS = ["ai_stance", "ai_uncertain", "ai_reason", "verified"]
AI_FILL = PatternFill("solid", fgColor="EDEDED")
AI_DOUBT_FILL = PatternFill("solid", fgColor="F8CBAD")
# Values that say nothing about trust in science.
NOT_APPLICABLE = "not_applicable"


def stance_values(codebook: Codebook) -> list[str]:
    return list(codebook.variable(STANCE).values)


def workbook_path(phase: str, annotator: str, sheets_dir: Path = SHEETS_DIR) -> Path:
    return sheets_dir / f"stance_{phase}_{annotator}.xlsx"


def workbook_paths(phase: str, sheets_dir: Path = SHEETS_DIR) -> dict[str, Path]:
    """Annotator ID -> stance workbook of a phase."""
    prefix = f"stance_{phase}_"
    return {path.stem.removeprefix(prefix): path for path in sorted(sheets_dir.glob(f"{prefix}*.xlsx"))}


def label_path(phase: str, annotator: str, labels_dir: Path = LABELS_DIR) -> Path:
    return labels_dir / phase / f"stance_{annotator}.csv"


def write_workbook(path: Path, annotator: str, phase: str, units: pd.DataFrame, codebook: Codebook) -> None:
    """Write one annotator's stance workbook; `units` are Reddit units in presentation order."""
    variable = codebook.variable(STANCE)
    rule = next(rule.text for rule in codebook.rules if rule.code == "R14")
    workbook = Workbook()
    workbook.properties.title = f"stance_{phase}_{annotator}"
    workbook.properties.description = json.dumps(
        {"phase": phase, "annotator": annotator, "codebook_version": codebook.version, "pass": "stance"}
    )
    instructions = workbook.active
    instructions.title = "Instrukcja"
    rows = [
        [f"Arkusz stanowiska: {phase}, annotator {annotator}"],
        [f"Codebook v{codebook.version} z {codebook.date.isoformat()}. Pełna instrukcja: codebook/codebook.md i docs/instrukcja_annotacji_reddit.md"],
        [],
        ["Zasady pracy"],
        ["1. Koduj samodzielnie. Nie porównuj etykiet z innymi przed zakończeniem całego arkusza."],
        ["2. Opieraj się tylko na tytule, treści i adresie linku w wierszu. Nie otwieraj linków."],
        [f"3. {STANCE}: jedna wartość z tabeli poniżej (lista rozwijana)."],
        ["4. W razie wątpliwości wybierz najlepszą wartość, ustaw uncertain = 1 i opisz problem w notes (R12)."],
        ["5. Nie zmieniaj kolumn unit_id ani kolejności wierszy."],
        [],
        [variable.question],
        ["Wartość", "Znaczenie"],
        *([value, meaning] for value, meaning in variable.values.items()),
        [],
        ["R14", rule],
    ]
    for row in rows:
        instructions.append(row)
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column == 1 and cell.value in {"Zasady pracy", "Wartość"}:
                for header_cell in row:
                    header_cell.font = HEADER_FONT
    instructions["A1"].font = Font(bold=True, size=14)
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 110

    sheet = workbook.create_sheet(SHEET_NAME)
    content = CONTENT_COLUMNS["reddit"]
    sheet.append([header for _, header, _ in content] + LABEL_COLUMNS)
    for unit in units.itertuples(index=False):
        values = [str(getattr(unit, column))[:EXCEL_CELL_LIMIT] for column, _, _ in content]
        sheet.append(values + [None] * (len(LABEL_COLUMNS) - 1) + [annotator])
    last_row = max(len(units) + 1, 2)
    for index, (_, _, width) in enumerate(content, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offset, name in enumerate(LABEL_COLUMNS, start=len(content) + 1):
        letter = get_column_letter(offset)
        sheet.column_dimensions[letter].width = 30 if name == "notes" else 16
        for row in range(2, len(units) + 2):
            sheet.cell(row, offset).fill = LABEL_FILL
        if name == STANCE:
            validation = DataValidation(type="list", formula1=f'"{",".join(variable.values)}"', allow_blank=True, showErrorMessage=True)
        elif name == "uncertain":
            validation = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True)
        else:
            continue
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{last_row}")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = HEADER_FONT
    sheet.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_workbook(path: Path, codebook: Codebook, partial: bool = False) -> tuple[pd.DataFrame, list[str]]:
    """Stance labels of a filled workbook and every problem found; with `partial`, rows without labels are skipped."""
    allowed = set(stance_values(codebook))
    workbook = load_workbook(path, read_only=True, data_only=True)
    metadata = _metadata(workbook)
    if SHEET_NAME not in workbook.sheetnames:
        return pd.DataFrame(columns=FILE_COLUMNS), [f"{path.name}: brak zakładki {SHEET_NAME}"]
    rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    header, body = list(rows[0]), rows[1:]
    expected = [name for _, name, _ in CONTENT_COLUMNS["reddit"]] + LABEL_COLUMNS
    if header[: len(expected)] != expected:
        return pd.DataFrame(columns=FILE_COLUMNS), [f"{SHEET_NAME}: nagłówki {header} różnią się od oczekiwanych {expected}."]
    labels, problems, unverified = [], [], []
    for excel_row, values in enumerate(body, start=2):
        record = dict(zip(header, values))
        if record["unit_id"] is None:
            continue
        suggestion = str(record.get("ai_stance") or record.get(f"ai_{STANCE}") or "").strip()
        if suggestion:
            _, problem = _binary(record.get("verified"), required=False)
            if problem:
                problems.append(f"{SHEET_NAME}, wiersz {excel_row} ({record['unit_id']}): verified {problem}")
            if is_unverified(record):
                # A prefilled label nobody has checked is the AI's, not the person's.
                unverified.append(str(record["unit_id"]))
                continue
        filled = any(record[name] not in (None, "") for name in (STANCE, "uncertain", "notes"))
        if partial and not filled:
            continue
        where = f"{SHEET_NAME}, wiersz {excel_row} ({record['unit_id']})"
        stance = str(record[STANCE] or "").strip().lower()
        if not stance:
            problems.append(f"{where}: {STANCE} brak wartości")
            stance = None
        elif stance not in allowed:
            problems.append(f"{where}: {STANCE} ma wartość {stance!r} spoza {sorted(allowed)}")
            stance = None
        uncertain, problem = _binary(record["uncertain"], required=False)
        if problem:
            problems.append(f"{where}: uncertain {problem}")
        labels.append({
            "unit_id": str(record["unit_id"]),
            "platform": "reddit",
            "annotator": record["annotator"],
            "codebook_version": metadata.get("codebook_version", ""),
            STANCE: stance,
            "uncertain": 0 if pd.isna(uncertain) else uncertain,
            "notes": "" if record["notes"] is None else str(record["notes"]),
            "ai_stance": suggestion,
            "decided_by": str(record.get("decided_by") or "human"),
            "p_review": 1.0 if record.get("p_review") in (None, "") else float(record["p_review"]),
            "verified_by": sheet_verified_by(record),
        })
    if unverified and not partial:
        problems.append(f"{len(unverified)} wierszy z podpowiedzią AI bez verified = 1 (sprawdź etykietę i wpisz 1), "
                        f"np. {unverified[:3]}")
    return pd.DataFrame(labels, columns=FILE_COLUMNS), problems


def add_ai_suggestions(path: Path, suggestions: pd.DataFrame, source: str) -> int:
    """Prefill consensus_stance with AI suggestions, keeping each suggestion and its reason in read-only columns.

    `suggestions` has unit_id, ai_stance, ai_uncertain and ai_reason. A label the
    person has already entered stays. Returns the number of rows that got a suggestion.
    """
    if (path.parent / f"~${path.name}").exists():
        raise RuntimeError(f"{path.name} is open in Excel (lock file ~${path.name}); close it first.")
    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    header = [cell.value for cell in sheet[1]]
    expected = [name for _, name, _ in CONTENT_COLUMNS["reddit"]] + LABEL_COLUMNS
    if header[: len(expected)] != expected:
        raise ValueError(f"{path.name}: headers {header} differ from {expected}")
    if len(header) > len(expected):
        raise ValueError(f"{path.name} already has columns after {expected[-1]}: {header[len(expected):]}")
    by_unit = suggestions.astype({"unit_id": str}).set_index("unit_id")
    first = len(header) + 1
    for offset, name in enumerate(AI_COLUMNS):
        cell = sheet.cell(1, first + offset, name)
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(first + offset)].width = {"ai_reason": 50, "verified": 10}.get(name, 14)
    stance_column = header.index(STANCE) + 1
    filled = 0
    for row in range(2, sheet.max_row + 1):
        unit_id = sheet.cell(row, 1).value
        if unit_id is None or str(unit_id) not in by_unit.index:
            continue
        suggestion = by_unit.loc[str(unit_id)]
        if sheet.cell(row, stance_column).value in (None, ""):
            sheet.cell(row, stance_column, suggestion["ai_stance"])
        doubt = int(suggestion["ai_uncertain"]) == 1
        for offset, value in enumerate([suggestion["ai_stance"], int(suggestion["ai_uncertain"]), suggestion["ai_reason"], None]):
            cell = sheet.cell(row, first + offset, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = LABEL_FILL if AI_COLUMNS[offset] == "verified" else AI_DOUBT_FILL if doubt else AI_FILL
        filled += 1
    letter = get_column_letter(first + AI_COLUMNS.index("verified"))
    validation = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True)
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{max(sheet.max_row, 2)}")

    instructions = workbook["Instrukcja"]
    for line in [
        [],
        ["Podpowiedzi AI"],
        [f"Kolumnę {STANCE} wstępnie wypełnił asystent AI ({source}) na podstawie tytułu, treści i adresu linku."],
        ["ai_stance, ai_uncertain i ai_reason to jego podpowiedź i uzasadnienie: nie zmieniaj ich (porównujemy je z Twoją etykietą)."],
        [f"Sprawdź każdy wiersz według R14. Jeśli się nie zgadzasz, popraw {STANCE}. Własne wątpliwości: uncertain i notes."],
        ["Po sprawdzeniu wiersza wpisz verified = 1. Wiersze bez verified = 1 import pomija albo odrzuca."],
        ["Wiersze z ai_uncertain = 1 (pomarańczowe) budziły wątpliwość AI: sprawdź je szczególnie uważnie."],
    ]:
        instructions.append(line)
        if line and line[0] == "Podpowiedzi AI":
            instructions.cell(instructions.max_row, 1).font = HEADER_FONT
    metadata = json.loads(workbook.properties.description or "{}")
    workbook.properties.description = json.dumps({**metadata, "ai_suggestions": source})
    workbook.save(path)
    return filled


def ai_review(labels: pd.DataFrame) -> pd.DataFrame:
    """Per annotator: AI suggestions checked and how many the person changed.

    Rows decided by a rule (`decided_by = auto`) carry the rule label in their AI column, so they are left out.
    """
    by_person = labels["decided_by"].fillna("human").ne("auto") if "decided_by" in labels else True
    checked = labels[labels["ai_stance"].fillna("").ne("") & labels[STANCE].notna() & by_person]
    rows = [{"annotator": annotator, "checked": len(part), "changed": int(part[STANCE].ne(part["ai_stance"]).sum())}
            for annotator, part in checked.groupby("annotator")]
    table = pd.DataFrame(rows, columns=["annotator", "checked", "changed"])
    return table.assign(kept_share=1 - table["changed"] / table["checked"])


def save_labels(labels: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    labels.reindex(columns=FILE_COLUMNS).to_csv(path, index=False, encoding="utf-8")


def load_labels(phase: str, labels_dir: Path = LABELS_DIR) -> pd.DataFrame:
    """All annotators' stance labels of a phase (empty if none imported yet)."""
    paths = sorted((labels_dir / phase).glob("stance_*.csv"))
    if not paths:
        return pd.DataFrame(columns=FILE_COLUMNS)
    labels = pd.concat(
        [pd.read_csv(path, dtype={"unit_id": str, "annotator": str, "codebook_version": str, STANCE: str, "notes": str,
                                  "ai_stance": str},
                     keep_default_na=False, na_values={STANCE: [""]}) for path in paths],
        ignore_index=True,
    )
    return labels.reindex(columns=FILE_COLUMNS).astype({"uncertain": "Int64"})


def reference(labels: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """One stance per unit: kept when everyone who coded the unit agrees; the rest waits for adjudication."""
    labels = labels[labels[STANCE].notna()]
    values_per_unit = labels.groupby("unit_id")[STANCE].nunique()
    agreed = values_per_unit[values_per_unit.eq(1)].index
    kept = labels[labels["unit_id"].isin(agreed)].drop_duplicates("unit_id").reset_index(drop=True)
    return kept, labels[~labels["unit_id"].isin(agreed)].reset_index(drop=True)


def agreement(labels: pd.DataFrame, samples: int = 2000, seed: int = 20261024) -> pd.DataFrame:
    """Agreement on units coded by two or more people: observed share, Cohen's kappa (two people) and Krippendorff's alpha."""
    labels = labels[labels[STANCE].notna()]
    matrix = labels.pivot_table(index="unit_id", columns="annotator", values=STANCE, aggfunc="first")
    rows = []
    for variable, transform in [(STANCE, lambda value: value), ("trust_relevant", lambda value: value != NOT_APPLICABLE)]:
        coded = matrix.map(lambda value: None if pd.isna(value) else transform(value))
        # Missing labels come back as NaN from DataFrame.map, so test with pd.isna, not `is None`.
        units = [[value for value in row if not pd.isna(value)] for row in coded.itertuples(index=False)]
        units = [row for row in units if len(row) >= 2]
        kappa = (
            cohen_kappa(coded.iloc[:, 0].tolist(), coded.iloc[:, 1].tolist(), samples, seed)
            if coded.shape[1] == 2 else None
        )
        alpha = krippendorff_alpha(units, samples=samples, seed=seed)
        rows.append({
            "variable": variable,
            "units": len(units),
            "annotators": coded.shape[1],
            "observed_agreement": float(np.mean([len(set(row)) == 1 for row in units])) if units else np.nan,
            "kappa": kappa.value if kappa else np.nan,
            "kappa_low": kappa.low if kappa else np.nan,
            "kappa_high": kappa.high if kappa else np.nan,
            "alpha": alpha.value,
            "alpha_low": alpha.low,
            "alpha_high": alpha.high,
        })
    return pd.DataFrame(rows)


ENRICHED_PHASE = "enriched"


def validation_frame(gold_labels: pd.DataFrame, manifest: pd.DataFrame, split: str,
                     enriched_labels: pd.DataFrame | None = None, enriched_sample: pd.DataFrame | None = None) -> pd.DataFrame:
    """Reference stances of one split with weights: the gold sample and, when given, the enriched r/conspiracy sample.

    Gold posts weigh N_h / coded (design weights); enriched posts 1 / inclusion
    probability. Each source's weights are rescaled to sum to its number of
    posts, as in the pooled audit (docs/hipotezy.md, 25.10.2026). `review_weight`
    is the weight divided by p_review for rows a person decided, and missing for
    rows accepted automatically: the main validity gate uses the reviewed rows only.
    """
    from digital_media_analysis.validation import design_weights

    parts = []
    gold, _ = reference(gold_labels)
    if len(gold):
        gold = gold.merge(design_weights(gold[["unit_id", "platform"]], manifest), on=["unit_id", "platform"])
        parts.append(gold[gold["split"].eq(split)].assign(source="gold"))
    if enriched_labels is not None and enriched_sample is not None and len(enriched_labels):
        enriched, _ = reference(enriched_labels)
        sample = enriched_sample[["unit_id", "split", "inclusion"]].astype({"unit_id": str})
        enriched = enriched.merge(sample, on="unit_id").assign(weight=lambda frame: 1 / frame["inclusion"])
        parts.append(enriched[enriched["split"].eq(split)].assign(source=ENRICHED_PHASE))
    columns = ["unit_id", "source", STANCE, "weight", "decided_by", "p_review", "review_weight"]
    if not parts:
        return pd.DataFrame(columns=columns)
    frames = []
    for part in parts:
        part = part.assign(weight=part["weight"] * len(part) / part["weight"].sum())
        decided = part["decided_by"].fillna("human").astype(str) if "decided_by" in part else pd.Series("human", index=part.index)
        p_review = pd.to_numeric(part["p_review"], errors="coerce").fillna(1.0) if "p_review" in part else 1.0
        frames.append(part.assign(decided_by=decided, p_review=p_review,
                                  review_weight=(part["weight"] / p_review).where(decided.eq("human"))))
    return pd.concat(frames, ignore_index=True)[columns]


# Plan B enrichment (docs/hipotezy.md, 26.10.2026): words that make a post more likely to state a stance on science.
SCIENCE_TERMS = re.compile(
    r"(?<!\w)(?:scien\w*|stud(?:y|ies)|research\w*|scholar\w*|doctor\w*|physician\w*|medic\w*|cancer\w*|tumou?r\w*|vaccin\w*"
    r"|vax\w*|virus\w*|viral|disease\w*|pandemic\w*|epidemic\w*|flu|chemtrail\w*|geoengineer\w*|autism|nasa|evolution\w*"
    r"|climate|warming|gmo\w*|monsanto|fluorid\w*|radiation|radioactiv\w*|nuclear|health\w*|pharma\w*|cure\w*|drug\w*|fda"
    r"|cdc|who|nutrition\w*|diet\w*|gene\w*|dna)(?!\w)",
    re.IGNORECASE,
)
ENRICHED_STRATA = {"topic": None, "science_terms": 150, "rest": 60}


def draw_enriched(units: pd.DataFrame, topic_codes: pd.Series, sizes: dict[str, int | None] = ENRICHED_STRATA,
                  annotators: Sequence[str] = ("IS", "JP"), seed: int = 20261026) -> pd.DataFrame:
    """Stratified sample of posts for the stance reference, with inclusion probabilities, splits and annotators.

    Strata: a consensus topic found by topics_v1 (`topic_codes` != none), else a
    science term (`SCIENCE_TERMS`), else the rest; None takes the whole stratum.
    Within each stratum half the posts go to dev and half to test, and the
    annotators alternate in a random order.
    """
    text = units["title"].fillna("") + " " + units["text"].fillna("")
    stratum = np.where(topic_codes.to_numpy() != "none", "topic", np.where(text.str.contains(SCIENCE_TERMS), "science_terms", "rest"))
    frame = units.assign(stratum=stratum)
    rng = np.random.default_rng(seed)
    parts = []
    for name, size in sizes.items():
        group = frame[frame["stratum"].eq(name)]
        chosen = group if size is None or len(group) <= size else group.sample(n=size, random_state=rng)
        order = rng.permutation(len(chosen))
        chosen = chosen.iloc[order].assign(
            stratum_size=len(group), inclusion=len(chosen) / len(group) if len(group) else np.nan,
            split=np.where(np.arange(len(chosen)) % 2 == 0, "dev", "test"),
        )
        parts.append(chosen)
    sample = pd.concat(parts, ignore_index=True)
    sample = sample.iloc[rng.permutation(len(sample))].reset_index(drop=True)
    return sample.assign(order=np.arange(1, len(sample) + 1),
                         annotator=[annotators[index % len(annotators)] for index in range(len(sample))])
