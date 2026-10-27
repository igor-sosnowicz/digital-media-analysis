"""AI suggestions in annotation workbooks, which a person checks before they count.

`prefill` writes an AI's labels into the person's columns of a sheet (only empty
cells) and keeps the suggestion, its doubt and its reason in columns `ai_*`
appended to the right, with a column `verified` for the person. Readers skip a
row that carries a suggestion until it has `verified = 1` (`unverified`), so a
label nobody checked never enters the reference data. The suggestion stays
next to the final label, so the share of changed suggestions can be reported.
"""

import json
from collections.abc import Sequence
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# `unverified` lives with the workbook readers that use it; it is re-exported here.
from digital_media_analysis.annotation_sheets import HEADER_FONT, LABEL_FILL, VERIFIED, unverified  # noqa: F401

AI_FILL = PatternFill("solid", fgColor="EDEDED")
AI_DOUBT_FILL = PatternFill("solid", fgColor="F8CBAD")
INSTRUCTIONS = [
    "Podpowiedzi AI",
    "Żółte kolumny z etykietami wstępnie wypełnił asystent AI ({source}).",
    "Kolumny ai_* to jego podpowiedź, wątpliwość (ai_uncertain) i uzasadnienie (ai_reason): nie zmieniaj ich.",
    "Sprawdź każdy wiersz według codebooka. Jeśli się nie zgadzasz, popraw etykietę; własne wątpliwości wpisz w uncertain i notes.",
    "Po sprawdzeniu wiersza wpisz verified = 1. Wiersze bez verified = 1 import pomija albo odrzuca.",
    "Wiersze z ai_uncertain = 1 (pomarańczowe) budziły wątpliwość AI: sprawdź je szczególnie uważnie.",
]


def ai_column(target: str) -> str:
    return f"ai_{target}"


def prefill(path: Path, sheet_name: str, key_columns: Sequence[str], suggestions: pd.DataFrame,
            targets: Sequence[str], source: str, instructions_sheet: str | None = "Instrukcja") -> int:
    """Fill `targets` of a sheet from `suggestions` and append the AI columns; returns the rows that got a suggestion.

    `suggestions` has `key_columns`, `ai_<target>` for each target, `ai_uncertain`
    and `ai_reason`. A cell the person already filled stays as it is.
    """
    if (path.parent / f"~${path.name}").exists():
        raise RuntimeError(f"{path.name} is open in Excel (lock file ~${path.name}); close it first.")
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    header = [cell.value for cell in sheet[1]]
    added = [*(ai_column(target) for target in targets), "ai_uncertain", "ai_reason", VERIFIED]
    if clash := [name for name in added if name in header]:
        raise ValueError(f"{path.name}/{sheet_name} already has AI columns {clash}")
    if missing := [name for name in [*key_columns, *targets] if name not in header]:
        raise ValueError(f"{path.name}/{sheet_name} lacks columns {missing}")
    by_key = {tuple(str(value) for value in row[list(key_columns)]): row for _, row in suggestions.iterrows()}
    first = len(header) + 1
    for offset, name in enumerate(added):
        sheet.cell(1, first + offset, name).font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(first + offset)].width = {"ai_reason": 50, VERIFIED: 10}.get(name, 16)
    filled = 0
    for row in range(2, sheet.max_row + 1):
        key = tuple(str(sheet.cell(row, header.index(column) + 1).value) for column in key_columns)
        if key not in by_key:
            continue
        suggestion = by_key[key]
        for target in targets:
            cell = sheet.cell(row, header.index(target) + 1)
            value = suggestion[ai_column(target)]
            if cell.value in (None, "") and not pd.isna(value):
                cell.value = value
        doubt = int(suggestion["ai_uncertain"]) == 1
        values = [*(suggestion[ai_column(target)] for target in targets), int(suggestion["ai_uncertain"]), suggestion["ai_reason"], None]
        for offset, value in enumerate(values):
            cell = sheet.cell(row, first + offset, None if pd.isna(value) else value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = LABEL_FILL if added[offset] == VERIFIED else AI_DOUBT_FILL if doubt else AI_FILL
        filled += 1
    letter = get_column_letter(first + len(added) - 1)
    validation = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True)
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{max(sheet.max_row, 2)}")

    if instructions_sheet and instructions_sheet in workbook.sheetnames:
        notes = workbook[instructions_sheet]
        if not any(cell.value == INSTRUCTIONS[0] for cell in notes["A"]):
            notes.append([])
            for line in INSTRUCTIONS:
                notes.append([line.format(source=source)])
                if line == INSTRUCTIONS[0]:
                    notes.cell(notes.max_row, 1).font = HEADER_FONT
    try:
        metadata = json.loads(workbook.properties.description or "{}")
    except json.JSONDecodeError:
        metadata = {}
    workbook.properties.description = json.dumps({**metadata, "ai_suggestions": source})
    workbook.save(path)
    return filled



def fill_queue_rows(path: Path, sheet: str | None, key_columns: Sequence[str], suggestions: pd.DataFrame, target: str) -> int:
    """Put AI suggestions into the rows a person decides (`decided_by = human`) of a review sheet; returns rows filled.

    The target cell gets the suggestion where it is empty, and the columns
    ai_<target>, ai_uncertain and ai_reason hold the suggestion, the AI's doubt and
    its reason; `verified` stays empty, so a row counts only after a person checks
    it. Rows decided by a rule are left as they are. Works on workbooks and CSV files.
    """
    keys = list(key_columns)
    by_key = {tuple(str(value) for value in row[keys]): row for _, row in suggestions.iterrows()}
    columns = [target, f"ai_{target}", "ai_uncertain", "ai_reason"]
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        filled = 0
        for index, row in frame.iterrows():
            key = tuple(str(row[column]) for column in keys)
            if row.get("decided_by") != "human" or key not in by_key:
                continue
            suggestion = by_key[key]
            if not str(row[target]).strip():
                frame.at[index, target] = suggestion[f"ai_{target}"]
            frame.loc[index, columns[1:]] = [suggestion[f"ai_{target}"], str(suggestion["ai_uncertain"]), suggestion["ai_reason"]]
            filled += 1
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        return filled
    if (path.parent / f"~${path.name}").exists():
        raise RuntimeError(f"{path.name} is open in Excel (lock file ~${path.name}); close it first.")
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1]]
    position = {name: header.index(name) + 1 for name in [*keys, *columns, "decided_by"]}
    filled = 0
    for row in range(2, worksheet.max_row + 1):
        key = tuple(str(worksheet.cell(row, position[column]).value) for column in keys)
        if worksheet.cell(row, position["decided_by"]).value != "human" or key not in by_key:
            continue
        suggestion = by_key[key]
        value = suggestion[f"ai_{target}"]
        value = None if pd.isna(value) else value
        if worksheet.cell(row, position[target]).value in (None, ""):
            worksheet.cell(row, position[target], value)
        worksheet.cell(row, position[f"ai_{target}"], value)
        worksheet.cell(row, position["ai_uncertain"], int(suggestion["ai_uncertain"]))
        worksheet.cell(row, position["ai_reason"], suggestion["ai_reason"])
        filled += 1
    workbook.save(path)
    return filled
